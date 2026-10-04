"""Main window of PDF QA Viewer."""
from __future__ import annotations

import csv
import json
import os
from datetime import datetime
from collections import Counter
from contextlib import contextmanager

import pymupdf as fitz
from PySide6.QtCore import Qt, QRectF, QPointF, QTimer
from PySide6.QtGui import QColor, QFont, QKeySequence, QPixmap, QShortcut, QBrush
from PySide6.QtWidgets import (
    QAbstractItemView, QApplication, QButtonGroup, QCheckBox, QColorDialog, QComboBox,
    QDialog, QDialogButtonBox, QDockWidget, QDoubleSpinBox, QFileDialog, QFormLayout,
    QGridLayout, QGroupBox, QHBoxLayout, QHeaderView, QLabel, QLineEdit, QMainWindow,
    QMessageBox, QPlainTextEdit, QProgressDialog, QPushButton, QRadioButton, QSizePolicy,
    QSpinBox, QSplitter, QTableWidget, QTableWidgetItem, QTabWidget, QToolBar, QToolButton,
    QSlider, QInputDialog, QMenu,
    QTreeWidget, QTreeWidgetItem, QVBoxLayout, QWidget,
)

from . import core
from .viewer import PdfView, np_to_qimage

# --------------------------------------------------------------------------- #
RED = QColor(235, 64, 52)
GREEN = QColor(32, 170, 90)
ORANGE = QColor(255, 159, 28)
BLUE = QColor("#4f8cff")
CYAN = QColor("#22d3ee")
YELLOW = QColor("#ffcc00")
PURPLE = QColor("#b072ff")
GREY = QColor("#7d8799")

KIND_COL = {
    "geo_removed": RED, "removed": RED,
    "geo_added": GREEN, "added": GREEN,
    "geo_changed": ORANGE, "changed": ORANGE,
}
ZONE_COL = {"": PURPLE, "OK": GREEN, "CHANGED": ORANGE, "FAIL": RED, "N/A": GREY, "DATA": BLUE}
ROW_BG = {
    "equal": None,
    "changed": QColor(255, 159, 28, 55),
    "removed": QColor(235, 64, 52, 60),
    "added": QColor(32, 170, 90, 60),
}
ROW_SYM = {"equal": "=", "changed": "≠", "removed": "−", "added": "+"}
TOOL_COL = {"pan": BLUE, "crop": CYAN, "read": YELLOW, "zone": PURPLE, "highlight": QColor("#ffe600")}
HL_SWATCHES = [("#ffe600", "Vàng"), ("#7CFC00", "Xanh lá"), ("#00e5ff", "Xanh dương"),
               ("#ff6ec7", "Hồng"), ("#ff9f1c", "Cam"), ("#ff3b30", "Đỏ")]
TOOL_HINT = {
    "pan": "Pan: kéo chuột trái để di chuyển, lăn chuột để zoom.",
    "crop": "Crop: kéo khung vùng cần cắt → xuất PDF vector / PNG.",
    "read": "Đọc data: kéo khung vùng cần đọc → so sánh text Ver1/Ver2 + tô màu khác biệt.",
    "zone": "Vùng check: kéo khung để định nghĩa vùng kiểm tra (mini zone).",
    "highlight": "Highlight: kéo qua text để tô theo dòng chữ, kéo vùng trống để tô vùng. "
                 "Chuột phải vào highlight để ghi chú / đổi màu / xoá. Ctrl+Z hoàn tác.",
}

QSS = """
* { font-family: 'Segoe UI', 'Inter', sans-serif; font-size: 10pt; color: #d7dce5; }
QMainWindow, QDialog { background: #12151c; }
QWidget { background: transparent; }
QDockWidget > QWidget, QTabWidget > QWidget, #panel { background: #171b24; }
QTabWidget { background: #171b24; }
QAbstractScrollArea::corner { background: #12151c; }
QToolBar { background: qlineargradient(x1:0,y1:0,x2:0,y2:1, stop:0 #1d2230, stop:1 #181c26);
           border: none; border-bottom: 1px solid #262c3a; padding: 6px 8px; spacing: 4px; }
QToolBar::separator { background: #2c3342; width: 1px; margin: 4px 8px; }
QLabel#logo { font-size: 13pt; font-weight: 700; color: white; padding: 0 10px 0 4px; }
QLabel#muted { color: #8b95a7; }
QLabel#summary { background: #1d2230; border: 1px solid #2a3140; border-radius: 8px; padding: 8px; }
QToolButton { border: 1px solid transparent; border-radius: 8px; padding: 6px 10px; color: #cfd5e0; }
QToolButton:hover { background: #262d3b; border-color: #343d4f; }
QToolButton:checked { background: qlineargradient(x1:0,y1:0,x2:1,y2:0, stop:0 #2b4a8a, stop:1 #3b3a8f);
                      border-color: #5b8dff; color: white; }
QToolButton#primary { background: qlineargradient(x1:0,y1:0,x2:1,y2:0, stop:0 #4f8cff, stop:1 #8a5cff);
                      color: white; font-weight: 600; padding: 6px 14px; }
QToolButton#primary:hover { background: qlineargradient(x1:0,y1:0,x2:1,y2:0, stop:0 #6a9dff, stop:1 #9d76ff); }
QPushButton { background: #232a37; border: 1px solid #323b4c; border-radius: 8px; padding: 7px 12px; }
QPushButton:hover { background: #2a3343; border-color: #4f8cff; }
QPushButton:pressed { background: #1e2533; }
QPushButton:disabled { color: #5b6578; border-color: #262c38; }
QPushButton#primary { background: qlineargradient(x1:0,y1:0,x2:1,y2:0, stop:0 #4f8cff, stop:1 #8a5cff);
                      border: none; color: white; font-weight: 600; }
QPushButton#primary:hover { background: qlineargradient(x1:0,y1:0,x2:1,y2:0, stop:0 #6a9dff, stop:1 #9d76ff); }
QTabWidget::pane { border: none; border-top: 1px solid #262c3a; }
QTabWidget::tab-bar { left: 0; }
QTabBar { background: #171b24; qproperty-drawBase: 0; }
QTabBar::tab { background: transparent; padding: 9px 12px; color: #8b95a7; border: none;
               border-bottom: 2px solid transparent; font-weight: 600; }
QTabBar::tab:hover { color: #d7dce5; }
QTabBar::tab:selected { color: white; border-bottom: 2px solid #6a8dff; }
QTableWidget, QTreeWidget, QPlainTextEdit, QLineEdit, QComboBox, QSpinBox, QDoubleSpinBox {
    background: #1b202b; border: 1px solid #2a3140; border-radius: 6px; padding: 3px;
    selection-background-color: #2b4a8a; selection-color: white; }
QLineEdit:focus, QComboBox:focus, QSpinBox:focus, QDoubleSpinBox:focus { border-color: #4f8cff; }
QComboBox QAbstractItemView { background: #1b202b; border: 1px solid #2a3140; }
QTableWidget, QTreeWidget { gridline-color: #262c3a; alternate-background-color: #1e2430; }
QHeaderView::section { background: #1f2531; color: #8b95a7; border: none; border-right: 1px solid #262c3a;
                       padding: 6px; font-weight: 600; }
QGroupBox { border: 1px solid #2a3140; border-radius: 10px; margin-top: 16px; padding: 10px 8px 8px 8px; }
QGroupBox::title { subcontrol-origin: margin; left: 12px; padding: 0 6px; color: #8b95a7; font-weight: 600; }
QDockWidget { titlebar-close-icon: none; }
QDockWidget::title { background: #1b202b; padding: 8px; text-align: left; font-weight: 700; }
QStatusBar { background: #171b24; color: #8b95a7; border-top: 1px solid #262c3a; }
QStatusBar QLabel { color: #8b95a7; padding: 0 8px; }
QSplitter::handle { background: #262c3a; }
QScrollBar:vertical { background: transparent; width: 10px; margin: 0; }
QScrollBar::handle:vertical { background: #343c4d; border-radius: 5px; min-height: 30px; }
QScrollBar:horizontal { background: transparent; height: 10px; margin: 0; }
QScrollBar::handle:horizontal { background: #343c4d; border-radius: 5px; min-width: 30px; }
QScrollBar::add-line, QScrollBar::sub-line { width: 0; height: 0; }
QScrollBar::add-page, QScrollBar::sub-page { background: none; }
QCheckBox::indicator, QRadioButton::indicator { width: 15px; height: 15px; border: 1px solid #3f4a5e;
                                               background: #1b202b; }
QCheckBox::indicator { border-radius: 4px; }
QRadioButton::indicator { border-radius: 8px; }
QCheckBox::indicator:hover, QRadioButton::indicator:hover { border-color: #4f8cff; }
QCheckBox::indicator:checked { background: qlineargradient(x1:0,y1:0,x2:1,y2:1, stop:0 #4f8cff, stop:1 #8a5cff);
                               border-color: #6a8dff; }
QRadioButton::indicator:checked { background: qradialgradient(cx:0.5, cy:0.5, radius:0.5, fx:0.5, fy:0.5,
                                  stop:0 #ffffff, stop:0.35 #ffffff, stop:0.45 #4f8cff, stop:1 #4f8cff);
                                  border-color: #6a8dff; }
QCheckBox::indicator:disabled, QRadioButton::indicator:disabled { background: #151922; border-color: #262c38; }
QToolTip { background: #232a37; color: #e6e9ef; border: 1px solid #3a4456; padding: 5px; }
QProgressDialog { background: #171b24; }
"""


@contextmanager
def busy():
    QApplication.setOverrideCursor(Qt.WaitCursor)
    try:
        yield
    finally:
        QApplication.restoreOverrideCursor()


def qrect(t) -> QRectF:
    return QRectF(QPointF(t[0], t[1]), QPointF(t[2], t[3])).normalized()


def frect(r: QRectF) -> fitz.Rect:
    return fitz.Rect(r.left(), r.top(), r.right(), r.bottom())


def ttuple(r: QRectF) -> tuple:
    return (round(r.left(), 2), round(r.top(), 2), round(r.right(), 2), round(r.bottom(), 2))


# --------------------------------------------------------------------------- #
#  Dialogs
# --------------------------------------------------------------------------- #
class ZoneDialog(QDialog):
    HINTS = {
        "compare": "So sánh text + nét vẽ trong vùng giữa Ver1 và Ver2 → OK / CHANGED.",
        "extract": "Chỉ đọc & lưu text trong vùng (VD: số bản vẽ, rev, ngày) → xuất Excel. Dùng được với 1 bản.",
        "contains": "Text trong vùng (Ver2, nếu không có thì Ver1) phải chứa giá trị. VD: BACKDRAFTING",
        "equals": "Text trong vùng phải bằng chính xác giá trị.",
        "regex": r"Text trong vùng phải khớp biểu thức. VD: ^L\d+  hoặc  REV\s*[B-Z]",
        "not_empty": "Vùng phải có text (không được để trống).",
        "number_range": "Mọi số trong vùng phải nằm trong khoảng. VD: 100-500",
    }

    def __init__(self, parent, name: str, rule="compare", expected="", preview="", allow_compare=True):
        super().__init__(parent)
        self.setWindowTitle("Định nghĩa vùng check")
        self.setMinimumWidth(460)
        lay = QVBoxLayout(self)
        form = QFormLayout()
        self.ed_name = QLineEdit(name)
        self.cb_rule = QComboBox()
        for k, v in core.RULES.items():
            self.cb_rule.addItem(v, k)
        if not allow_compare:
            self.cb_rule.model().item(0).setEnabled(False)
            self.cb_rule.setItemText(0, core.RULES["compare"] + "  (cần mở cả 2 bản)")
            if rule == "compare":
                rule = "extract"
        self.cb_rule.setCurrentIndex(max(0, list(core.RULES).index(rule) if rule in core.RULES else 0))
        self.ed_exp = QLineEdit(expected)
        self.lbl_hint = QLabel()
        self.lbl_hint.setObjectName("muted")
        self.lbl_hint.setWordWrap(True)
        form.addRow("Tên vùng", self.ed_name)
        form.addRow("Quy tắc", self.cb_rule)
        form.addRow("Giá trị", self.ed_exp)
        form.addRow("", self.lbl_hint)
        lay.addLayout(form)
        if preview:
            lay.addWidget(QLabel("Text hiện có trong vùng:"))
            pv = QPlainTextEdit(preview)
            pv.setReadOnly(True)
            pv.setMaximumHeight(130)
            lay.addWidget(pv)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        lay.addWidget(bb)
        self.cb_rule.currentIndexChanged.connect(self._upd)
        self._upd()

    def _upd(self):
        k = self.cb_rule.currentData()
        self.lbl_hint.setText(self.HINTS.get(k, ""))
        self.ed_exp.setEnabled(k not in ("compare", "not_empty", "extract"))

    def values(self):
        return self.ed_name.text().strip() or "Zone", self.cb_rule.currentData(), self.ed_exp.text()


class CropDialog(QDialog):
    def __init__(self, parent, has1: bool, has2: bool):
        super().__init__(parent)
        self.setWindowTitle("Crop vùng chọn")
        self.setMinimumWidth(380)
        lay = QVBoxLayout(self)
        g1 = QGroupBox("Nguồn")
        v = QVBoxLayout(g1)
        self.chk1 = QCheckBox("Ver1")
        self.chk2 = QCheckBox("Ver2")
        self.chkd = QCheckBox("Ảnh overlay khác biệt (PNG)")
        self.chk1.setEnabled(has1)
        self.chk1.setChecked(has1)
        self.chk2.setEnabled(has2)
        self.chk2.setChecked(has2 and not has1)
        self.chkd.setEnabled(has1 and has2)
        for w in (self.chk1, self.chk2, self.chkd):
            v.addWidget(w)
        lay.addWidget(g1)
        g2 = QGroupBox("Định dạng")
        f = QFormLayout(g2)
        self.rb_pdf = QRadioButton("PDF vector (giữ nét & text)")
        self.rb_png = QRadioButton("Ảnh PNG")
        self.rb_pdf.setChecked(True)
        self.sp_dpi = QSpinBox()
        self.sp_dpi.setRange(72, 1200)
        self.sp_dpi.setValue(300)
        self.sp_dpi.setSuffix(" dpi")
        f.addRow(self.rb_pdf)
        f.addRow(self.rb_png)
        f.addRow("Độ phân giải ảnh", self.sp_dpi)
        lay.addWidget(g2)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        lay.addWidget(bb)


class BatchResultDialog(QDialog):
    def __init__(self, win: "MainWindow", results, names):
        super().__init__(win)
        self.win, self.results, self.names = win, results, names
        self.setWindowTitle("Kết quả kiểm tra nhiều file")
        self.resize(1100, 640)
        lay = QVBoxLayout(self)
        cnt = Counter(z.status for _, z in results)
        nfiles = len({p for p, _ in results})
        parts = [f"<b>{nfiles}</b> file · <b>{len(results)}</b> vùng"]
        for st in ("OK", "FAIL", "DATA", "N/A"):
            if cnt[st]:
                parts.append(f"<span style='color:{ZONE_COL[st].name()}'>{st}: <b>{cnt[st]}</b></span>")
        lb = QLabel(" &nbsp;·&nbsp; ".join(parts))
        lb.setObjectName("summary")
        lb.setTextFormat(Qt.RichText)
        lay.addWidget(lb)
        row = QHBoxLayout()
        self.cb_filter = QComboBox()
        self.cb_filter.addItems(["Tất cả", "Chỉ FAIL", "Chỉ OK", "Chỉ DATA", "Chỉ N/A"])
        self.cb_filter.currentIndexChanged.connect(self._fill)
        row.addWidget(QLabel("Lọc:"))
        row.addWidget(self.cb_filter)
        row.addStretch()
        hint = QLabel("Double-click 1 dòng để mở bản vẽ tại vùng đó")
        hint.setObjectName("muted")
        row.addWidget(hint)
        lay.addLayout(row)
        self.t = MainWindow._table(["File", "Trang", "Vùng", "Quy tắc", "Kết quả", "Text", "Chi tiết"], 5)
        for c, wd in enumerate([260, 50, 90, 120, 70, 300, 200]):
            self.t.setColumnWidth(c, wd)
        self.t.itemDoubleClicked.connect(self._open)
        lay.addWidget(self.t, 1)
        bb = QHBoxLayout()
        b1 = QPushButton("📊 Xuất Excel")
        b1.setObjectName("primary")
        b1.clicked.connect(self._export)
        b2 = QPushButton("Đóng")
        b2.clicked.connect(self.accept)
        bb.addStretch()
        bb.addWidget(b1)
        bb.addWidget(b2)
        lay.addLayout(bb)
        self._fill()

    def _fill(self):
        f = {0: None, 1: "FAIL", 2: "OK", 3: "DATA", 4: "N/A"}[self.cb_filter.currentIndex()]
        self.rows = [(p, z) for p, z in self.results if f is None or z.status == f]
        t = self.t
        t.setRowCount(0)
        for path, z in self.rows:
            r = t.rowCount()
            t.insertRow(r)
            txt = z.text_v1.replace("\n", " ⏎ ")
            vals = [os.path.basename(path), str(z.page + 1), z.name, core.RULES.get(z.rule, z.rule),
                    z.status, txt, z.detail]
            for c, v in enumerate(vals):
                it = QTableWidgetItem(v)
                it.setToolTip(path if c == 0 else v)
                if c == 4:
                    col = QColor(ZONE_COL.get(z.status, GREY))
                    it.setForeground(QBrush(col))
                    bg = QColor(col)
                    bg.setAlpha(45)
                    it.setBackground(bg)
                t.setItem(r, c, it)

    def _open(self, item):
        path, z = self.rows[item.row()]
        w = self.win
        if not (w.docs[1] and os.path.normcase(w.docs[1].path) == os.path.normcase(path)):
            w.open_doc(1, path)
        if w.mode != "side" or not w.docs[2]:
            w.set_mode("v1")
        w._goto(z.page, qrect(z.rect), 1, 1.6)

    def _export(self):
        path, _ = QFileDialog.getSaveFileName(self, "Xuất kết quả", os.path.join(self.win._last_dir, "batch_check.xlsx"),
                                              "Excel (*.xlsx)")
        if not path:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            fills = {"OK": "C6EFCE", "FAIL": "FFC7CE", "DATA": "DDEBF7", "N/A": "EDEDED"}
            wb = Workbook()
            ws = wb.active
            ws.title = "Chi tiet"
            ws.append(["File", "Trang", "Vùng", "Quy tắc", "Giá trị", "Kết quả", "Text", "Chi tiết", "Đường dẫn"])
            for path_, z in self.results:
                ws.append([os.path.basename(path_), z.page + 1, z.name, core.RULES.get(z.rule, z.rule), z.expected,
                           z.status, z.text_v1, z.detail, path_])
                if z.status in fills:
                    ws.cell(ws.max_row, 6).fill = PatternFill("solid", fgColor=fills[z.status])
            # data sheet: one row per file/page, one column per zone
            ws2 = wb.create_sheet("Bang data")
            ws2.append(["File", "Trang"] + self.names + ["Kiểm tra"])
            grid = {}
            for path_, z in self.results:
                grid.setdefault((path_, z.page), {})[z.name] = z
            for (path_, p), zs in grid.items():
                bad = [f"{n}:{zs[n].status}" for n in self.names if n in zs and zs[n].status == "FAIL"]
                ws2.append([os.path.basename(path_), p + 1]
                           + [zs[n].text_v1.replace("\n", " ") if n in zs else "" for n in self.names]
                           + ["; ".join(bad) or "OK"])
                ws2.cell(ws2.max_row, len(self.names) + 3).fill = PatternFill(
                    "solid", fgColor="FFC7CE" if bad else "C6EFCE")
            for sh in (ws, ws2):
                for c in sh[1]:
                    c.font = Font(bold=True, color="FFFFFF")
                    c.fill = PatternFill("solid", fgColor="2B4A8A")
                for col in sh.columns:
                    sh.column_dimensions[col[0].column_letter].width = min(
                        60, max(len(str(c.value or "")) for c in col[:300]) + 2)
                sh.freeze_panes = "A2"
            wb.save(path)
        except Exception as e:
            QMessageBox.critical(self, "Xuất", f"Lỗi: {e}")
            return
        if QMessageBox.question(self, "Xuất xong", f"Đã lưu:\n{path}\n\nMở file?") == QMessageBox.Yes:
            os.startfile(path)


# --------------------------------------------------------------------------- #
#  Main window
# --------------------------------------------------------------------------- #
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.docs: dict[int, core.PdfDoc | None] = {1: None, 2: None}
        self.page = 0
        self.mode = "v1"
        self.tool = "pan"
        self.zones: list[core.Zone] = []
        self.diff_items: dict[int, list[core.DiffItem]] = {}
        self.query_hits: list[core.QueryHit] = []
        self.query_color = QColor("#ff4fd8")
        self.highlights: list[core.Highlight] = []
        self._hl_id = 0
        self.hl_color = QColor(HL_SWATCHES[0][0])
        self.hl_sel: int | None = None
        self.region: tuple[int, QRectF] | None = None
        self.region_td: core.TextDiff | None = None
        self.view_src: dict[PdfView, object] = {}
        self._pd_key = None
        self._pd = None
        self._td_cache: dict = {}
        self._syncing = False
        self._last_dir = os.getcwd()

        self._hires_timer = QTimer(self, singleShot=True, interval=180)
        self._hires_timer.timeout.connect(self._do_hires)
        self._settings_timer = QTimer(self, singleShot=True, interval=350)
        self._settings_timer.timeout.connect(self._settings_applied)

        self._build_ui()
        self.set_mode("v1")
        self.set_tool("pan")

    # ================================================================== UI ==
    def _build_ui(self):
        self.setWindowTitle("PDF QA Viewer — So sánh & kiểm tra bản vẽ")
        self.resize(1680, 980)
        self.setAcceptDrops(True)

        self.viewA = PdfView("VER 1", BLUE.name())
        self.viewB = PdfView("VER 2", GREEN.name())
        for v in (self.viewA, self.viewB):
            v.rectSelected.connect(lambda r, v=v: self.on_rect(v, r))
            v.mouseMovedPdf.connect(self.on_mouse)
            v.contextAt.connect(lambda p, g, v=v: self.on_view_context(v, p, g))
            v.viewChanged.connect(lambda v=v: self.on_view_changed(v))
        self.splitter = QSplitter(Qt.Horizontal)
        self.splitter.addWidget(self.viewA)
        self.splitter.addWidget(self.viewB)
        self.splitter.setHandleWidth(3)
        self.setCentralWidget(self.splitter)

        self._build_toolbar()
        self._build_dock()
        self._build_status()
        self._build_shortcuts()

    def _tbtn(self, text, tip, slot=None, checkable=False, name=None):
        b = QToolButton()
        b.setText(text)
        b.setToolTip(tip)
        b.setCheckable(checkable)
        b.setCursor(Qt.PointingHandCursor)
        if name:
            b.setObjectName(name)
        if slot:
            b.clicked.connect(slot)
        self.tb.addWidget(b)
        return b

    def _build_toolbar(self):
        self.tb = QToolBar("Main")
        self.tb.setMovable(False)
        self.addToolBar(Qt.TopToolBarArea, self.tb)
        logo = QLabel("◆ PDF QA")
        logo.setObjectName("logo")
        self.tb.addWidget(logo)

        self._tbtn("📂 PDF 1", "Mở file PDF 1 (Ctrl+1)", lambda: self.open_doc(1))
        self._tbtn("📂 PDF 2", "Mở file PDF 2 (Ctrl+2)", lambda: self.open_doc(2))
        self.tb.addSeparator()

        self.mode_group = QButtonGroup(self)
        self.mode_btns = {}
        for key, text, tip in [("v1", "Ver1", "Chỉ xem Ver1"),
                               ("v2", "Ver2", "Chỉ xem Ver2"),
                               ("side", "⇆ Song song", "Xem Ver1 | Ver2 song song, đồng bộ zoom/pan"),
                               ("overlay", "◐ Overlay", "Chồng 2 bản: Đỏ = chỉ Ver1 (xoá), Xanh = chỉ Ver2 (thêm)")]:
            b = self._tbtn(text, tip, checkable=True)
            b.clicked.connect(lambda _=False, k=key: self.set_mode(k))
            self.mode_group.addButton(b)
            self.mode_btns[key] = b
        self.tb.addSeparator()

        self.tool_group = QButtonGroup(self)
        self.tool_btns = {}
        for key, text, tip in [("pan", "✋ Pan", "Di chuyển (H / Esc)"),
                               ("crop", "✂ Crop", "Kéo khung để crop PDF/PNG (C)"),
                               ("read", "🔎 Đọc data", "Kéo khung để đọc & so sánh data trong vùng (R)"),
                               ("zone", "▣ Vùng check", "Kéo khung để định nghĩa vùng check (Z)"),
                               ("highlight", "🖍 Highlight", "Tô highlight như Bluebeam: kéo qua text hoặc vùng (G)")]:
            b = self._tbtn(text, tip, checkable=True)
            b.clicked.connect(lambda _=False, k=key: self.set_tool(k))
            self.tool_group.addButton(b)
            self.tool_btns[key] = b
        self.tb.addSeparator()

        self._tbtn("◀", "Trang trước (PgUp)", lambda: self.set_page(self.page - 1))
        self.sp_page = QSpinBox()
        self.sp_page.setRange(1, 1)
        self.sp_page.setFixedWidth(64)
        self.sp_page.setAlignment(Qt.AlignCenter)
        self.sp_page.valueChanged.connect(lambda v: self.set_page(v - 1))
        self.tb.addWidget(self.sp_page)
        self.lbl_pages = QLabel(" / 0 ")
        self.lbl_pages.setObjectName("muted")
        self.tb.addWidget(self.lbl_pages)
        self._tbtn("▶", "Trang sau (PgDn)", lambda: self.set_page(self.page + 1))
        self.tb.addSeparator()
        self._tbtn("⤢ Fit", "Vừa màn hình (F)", self.fit)
        self._tbtn("−", "Thu nhỏ (-)", lambda: self._zoom(1 / 1.25))
        self._tbtn("+", "Phóng to (+)", lambda: self._zoom(1.25))

        sp = QWidget()
        sp.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
        self.tb.addWidget(sp)
        self._tbtn("⚡ So sánh trang", "So sánh Ver1 ↔ Ver2 trang hiện tại (Ctrl+D)",
                   self.compare_current, name="primary")

    def _build_dock(self):
        dock = QDockWidget("  Bảng QA", self)
        dock.setFeatures(QDockWidget.DockWidgetMovable | QDockWidget.DockWidgetFloatable)
        self.tabs = QTabWidget()
        self.tabs.addTab(self._tab_compare(), "⇆ So sánh")
        self.tabs.addTab(self._tab_read(), "🔎 Đọc data")
        self.tabs.addTab(self._tab_zones(), "▣ Vùng check")
        self.tabs.addTab(self._tab_query(), "⌕ Query")
        self.tabs.addTab(self._tab_markup(), "🖍 Markup")
        dock.setWidget(self.tabs)
        self.tabs.setDocumentMode(True)
        holder = QWidget()
        holder.setObjectName("panel")
        hl = QVBoxLayout(holder)
        hl.setContentsMargins(0, 0, 0, 0)
        hl.addWidget(self.tabs)
        dock.setWidget(holder)
        dock.setMinimumWidth(470)
        self.addDockWidget(Qt.RightDockWidgetArea, dock)
        self.resizeDocks([dock], [520], Qt.Horizontal)

    def _panel(self):
        w = QWidget()
        w.setObjectName("panel")
        lay = QVBoxLayout(w)
        lay.setContentsMargins(10, 10, 10, 10)
        lay.setSpacing(8)
        return w, lay

    @staticmethod
    def _table(headers, stretch_col=None):
        t = QTableWidget(0, len(headers))
        t.setHorizontalHeaderLabels(headers)
        t.verticalHeader().setVisible(False)
        t.setSelectionBehavior(QAbstractItemView.SelectRows)
        t.setEditTriggers(QAbstractItemView.NoEditTriggers)
        t.setAlternatingRowColors(True)
        t.setWordWrap(False)
        h = t.horizontalHeader()
        h.setSectionResizeMode(QHeaderView.Interactive)
        if stretch_col is not None:
            h.setSectionResizeMode(stretch_col, QHeaderView.Stretch)
        t.verticalHeader().setDefaultSectionSize(26)
        return t

    def _legend(self):
        w = QWidget()
        h = QHBoxLayout(w)
        h.setContentsMargins(0, 0, 0, 0)
        for col, txt in [(RED, "Chỉ có ở Ver1 (xoá)"), (GREEN, "Chỉ có ở Ver2 (thêm)"),
                         (ORANGE, "Sửa"), (GREY, "Không đổi")]:
            dot = QLabel("●")
            dot.setStyleSheet(f"color:{col.name()}; font-size:14pt;")
            lb = QLabel(txt)
            lb.setObjectName("muted")
            h.addWidget(dot)
            h.addWidget(lb)
            h.addSpacing(6)
        h.addStretch()
        return w

    def _tab_compare(self):
        w, lay = self._panel()
        grp = QGroupBox("Thiết lập so sánh")
        g = QGridLayout(grp)
        self.cb_dpi = QComboBox()
        self.cb_dpi.addItems(["Auto", "72", "100", "150", "200", "300"])
        self.cb_dpi.setToolTip("Độ phân giải render nền. Auto ≈ 5000px cạnh dài. Khi zoom sâu sẽ tự render nét.")
        self.sp_thr = QSpinBox()
        self.sp_thr.setRange(50, 254)
        self.sp_thr.setValue(200)
        self.sp_thr.setToolTip("Ngưỡng mực: pixel tối hơn giá trị này được coi là nét vẽ")
        self.sp_tol = QSpinBox()
        self.sp_tol.setRange(0, 6)
        self.sp_tol.setValue(1)
        self.sp_tol.setSuffix(" px")
        self.sp_tol.setToolTip("Dung sai lệch nét (bỏ qua lệch nhỏ do khử răng cưa)")
        self.sp_dx = QDoubleSpinBox()
        self.sp_dy = QDoubleSpinBox()
        for s in (self.sp_dx, self.sp_dy):
            s.setRange(-2000, 2000)
            s.setDecimals(1)
            s.setSingleStep(0.5)
            s.setSuffix(" pt")
            s.setToolTip("Lệch Ver2 so với Ver1 (nội dung tại Ver1 (x,y) nằm ở Ver2 (x+dx, y+dy))")
        self.chk_pixel = QCheckBox("So sánh nét vẽ (pixel)")
        self.chk_text = QCheckBox("So sánh text")
        self.chk_show = QCheckBox("Hiện khung khác biệt trên bản vẽ")
        for c in (self.chk_pixel, self.chk_text, self.chk_show):
            c.setChecked(True)
        g.addWidget(QLabel("DPI render"), 0, 0)
        g.addWidget(self.cb_dpi, 0, 1)
        g.addWidget(QLabel("Ngưỡng mực"), 0, 2)
        g.addWidget(self.sp_thr, 0, 3)
        g.addWidget(QLabel("Lệch Ver2 dx"), 1, 0)
        g.addWidget(self.sp_dx, 1, 1)
        g.addWidget(QLabel("dy"), 1, 2)
        g.addWidget(self.sp_dy, 1, 3)
        g.addWidget(QLabel("Dung sai"), 2, 0)
        g.addWidget(self.sp_tol, 2, 1)
        g.addWidget(self.chk_pixel, 3, 0, 1, 2)
        g.addWidget(self.chk_text, 3, 2, 1, 2)
        g.addWidget(self.chk_show, 4, 0, 1, 4)
        lay.addWidget(grp)
        for s in (self.sp_thr, self.sp_tol, self.sp_dx, self.sp_dy):
            s.valueChanged.connect(lambda _=None: self._settings_timer.start())
        self.cb_dpi.currentIndexChanged.connect(lambda _=None: self._settings_timer.start())
        self.chk_show.toggled.connect(lambda _=None: self.redraw_overlays())

        lay.addWidget(self._legend())
        row = QHBoxLayout()
        b1 = QPushButton("⚡ So sánh trang này")
        b1.setObjectName("primary")
        b1.clicked.connect(self.compare_current)
        b2 = QPushButton("📑 So sánh tất cả trang")
        b2.clicked.connect(self.compare_all)
        row.addWidget(b1)
        row.addWidget(b2)
        lay.addLayout(row)
        self.lbl_cmp = QLabel("Mở Ver1 và Ver2 rồi bấm So sánh.")
        self.lbl_cmp.setObjectName("summary")
        self.lbl_cmp.setWordWrap(True)
        self.lbl_cmp.setTextFormat(Qt.RichText)
        lay.addWidget(self.lbl_cmp)
        self.tree = QTreeWidget()
        self.tree.setHeaderLabels(["Loại", "Nội dung", "Vị trí (pt)"])
        self.tree.setAlternatingRowColors(True)
        self.tree.header().setSectionResizeMode(1, QHeaderView.Stretch)
        self.tree.setColumnWidth(0, 130)
        self.tree.itemClicked.connect(self.on_tree_click)
        lay.addWidget(self.tree, 1)
        b3 = QPushButton("📊 Xuất danh sách khác biệt (Excel)")
        b3.clicked.connect(self.export_diffs)
        lay.addWidget(b3)
        return w

    def _tab_read(self):
        w, lay = self._panel()
        info = QLabel("Chọn công cụ <b>🔎 Đọc data</b> (R) rồi kéo khung trên bản vẽ. "
                      "Text Ver1/Ver2 trong vùng được so sánh từng dòng và tô màu khác biệt.")
        info.setWordWrap(True)
        info.setObjectName("muted")
        lay.addWidget(info)
        self.lbl_read = QLabel("Chưa chọn vùng.")
        self.lbl_read.setObjectName("summary")
        self.lbl_read.setWordWrap(True)
        self.lbl_read.setTextFormat(Qt.RichText)
        lay.addWidget(self.lbl_read)
        self.read_preview = QLabel()
        self.read_preview.setAlignment(Qt.AlignCenter)
        self.read_preview.setMinimumHeight(200)
        self.read_preview.setMaximumHeight(260)
        self.read_preview.setStyleSheet("background:#0f1218; border:1px solid #2a3140; border-radius:8px;")
        lay.addWidget(self.read_preview)
        self.read_table = self._table(["", "Ver1", "Ver2"])
        self.read_table.setColumnWidth(0, 30)
        self.read_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
        self.read_table.horizontalHeader().setSectionResizeMode(2, QHeaderView.Stretch)
        lay.addWidget(self.read_table, 1)
        g = QGridLayout()
        btns = [("📋 Copy", self.read_copy), ("💾 Xuất CSV/Excel", self.read_export),
                ("▣ Thêm làm vùng check", lambda: self.region and self.add_zone(self.region[1])),
                ("✂ Crop vùng này", lambda: self.region and self.do_crop(self.region[1]))]
        for i, (t, s) in enumerate(btns):
            b = QPushButton(t)
            b.clicked.connect(s)
            g.addWidget(b, i // 2, i % 2)
        lay.addLayout(g)
        return w

    def _tab_zones(self):
        w, lay = self._panel()
        info = QLabel("Bấm <b>▣ Vùng check</b> (Z) rồi kéo khung để tạo vùng (VD: ô số bản vẽ, Rev, Suitability). "
                      "Vùng được áp cho mọi trang, hoặc cho nhiều file PDF cùng khổ giấy.")
        info.setWordWrap(True)
        info.setObjectName("muted")
        lay.addWidget(info)
        self.lbl_zone = QLabel("0 vùng")
        self.lbl_zone.setObjectName("summary")
        self.lbl_zone.setTextFormat(Qt.RichText)
        lay.addWidget(self.lbl_zone)
        self.zone_table = self._table(["Tên", "Trang", "Quy tắc", "Giá trị", "Kết quả",
                                       "Chi tiết", "Text Ver1", "Text Ver2"])
        for c, wd in enumerate([90, 50, 120, 90, 80, 160, 160, 160]):
            self.zone_table.setColumnWidth(c, wd)
        self.zone_table.setMinimumHeight(80)
        self.zone_table.itemClicked.connect(self.on_zone_click)
        self.zone_table.itemDoubleClicked.connect(lambda _: self.edit_zone())
        lay.addWidget(self.zone_table, 1)
        self.chk_allpages = QCheckBox("Áp dụng vùng cho tất cả trang của file")
        self.chk_allpages.setChecked(True)
        self.chk_allpages.setToolTip("Vùng vẽ ở 1 trang sẽ được kiểm tra ở cùng vị trí trên mọi trang")
        lay.addWidget(self.chk_allpages)
        row = QHBoxLayout()
        run = QPushButton("▶ Kiểm tra file đang mở")
        run.setObjectName("primary")
        run.clicked.connect(self.run_zone_check)
        batch = QPushButton("📁 Kiểm tra nhiều file PDF…")
        batch.setObjectName("primary")
        batch.setToolTip("Áp các vùng check cho mọi trang của nhiều file PDF (cùng khổ / cùng khung tên)")
        batch.clicked.connect(self.batch_check_files)
        row.addWidget(run)
        row.addWidget(batch)
        lay.addLayout(row)
        g = QGridLayout()
        g.setSpacing(6)
        btns = [("✎ Sửa", self.edit_zone), ("🗑 Xoá", self.delete_zones),
                ("💾 Lưu template", self.save_template), ("📂 Mở template", self.load_template),
                ("📊 Xuất Excel", self.export_zones), ("▦ Bảng data", self.export_zone_pivot),
                ("✂ Crop vùng", self.crop_zones)]
        for i, (t, s) in enumerate(btns):
            b = QPushButton(t)
            b.clicked.connect(s)
            g.addWidget(b, i // 4, i % 4)
        lay.addLayout(g)
        return w

    def _tab_query(self):
        w, lay = self._panel()
        row = QHBoxLayout()
        self.q_edit = QLineEdit()
        self.q_edit.setPlaceholderText("Nhập text / số / regex cần tìm, VD: 1200  hoặc  ^PT\\d+")
        self.q_edit.returnPressed.connect(self.run_query)
        b = QPushButton("⌕ Tìm")
        b.setObjectName("primary")
        b.clicked.connect(self.run_query)
        row.addWidget(self.q_edit, 1)
        row.addWidget(b)
        lay.addLayout(row)
        g = QGridLayout()
        self.chk_regex = QCheckBox("Regex")
        self.chk_case = QCheckBox("Phân biệt hoa/thường")
        self.chk_changed = QCheckBox("Chỉ kết quả có thay đổi Ver1↔Ver2")
        self.cb_scope = QComboBox()
        self.cb_scope.addItems(["Trang hiện tại", "Tất cả trang", "Trong các vùng check",
                                "Trong vùng Đọc data"])
        self.cb_ver = QComboBox()
        self.cb_ver.addItems(["Ver1 + Ver2", "Chỉ Ver1", "Chỉ Ver2"])
        self.btn_qcol = QPushButton("■ Màu highlight")
        self.btn_qcol.clicked.connect(self.pick_query_color)
        self._upd_qcol()
        g.addWidget(QLabel("Phạm vi"), 0, 0)
        g.addWidget(self.cb_scope, 0, 1)
        g.addWidget(QLabel("Bản"), 0, 2)
        g.addWidget(self.cb_ver, 0, 3)
        g.addWidget(self.chk_regex, 1, 0)
        g.addWidget(self.chk_case, 1, 1)
        g.addWidget(self.btn_qcol, 1, 2, 1, 2)
        g.addWidget(self.chk_changed, 2, 0, 1, 4)
        lay.addLayout(g)
        self.lbl_query = QLabel("Kết quả: —")
        self.lbl_query.setObjectName("summary")
        self.lbl_query.setTextFormat(Qt.RichText)
        lay.addWidget(self.lbl_query)
        self.q_table = self._table(["Bản", "Trang", "Text", "Trạng thái", "Vị trí (pt)"], 2)
        self.q_table.setColumnWidth(0, 50)
        self.q_table.setColumnWidth(1, 50)
        self.q_table.setColumnWidth(3, 110)
        self.q_table.itemClicked.connect(self.on_query_click)
        lay.addWidget(self.q_table, 1)
        row2 = QHBoxLayout()
        b1 = QPushButton("✕ Xoá highlight")
        b1.clicked.connect(self.clear_query)
        b2 = QPushButton("📊 Xuất Excel")
        b2.clicked.connect(self.export_query)
        row2.addWidget(b1)
        row2.addWidget(b2)
        lay.addLayout(row2)
        return w

    def _tab_markup(self):
        w, lay = self._panel()
        info = QLabel("Bấm <b>🖍 Highlight</b> (G) rồi kéo chuột: qua chữ → tô theo dòng chữ, "
                      "qua vùng trống → tô vùng. Chuột phải vào highlight để ghi chú / đổi màu / xoá.")
        info.setWordWrap(True)
        info.setObjectName("muted")
        lay.addWidget(info)
        grp = QGroupBox("Bút highlight")
        g = QGridLayout(grp)
        sw = QHBoxLayout()
        self.hl_group = QButtonGroup(self)
        for i, (hexc, name) in enumerate(HL_SWATCHES):
            b = QToolButton()
            b.setCheckable(True)
            b.setToolTip(name)
            b.setFixedSize(30, 26)
            b.setStyleSheet(f"QToolButton {{ background:{hexc}; border:2px solid #2a3140; border-radius:6px; }}"
                            f"QToolButton:checked {{ border:2px solid white; }}")
            b.clicked.connect(lambda _=False, c=hexc: self._set_hl_color(QColor(c)))
            self.hl_group.addButton(b)
            sw.addWidget(b)
            if i == 0:
                b.setChecked(True)
        more = QToolButton()
        more.setText("…")
        more.setToolTip("Màu khác")
        more.setFixedSize(30, 26)
        more.clicked.connect(self._pick_hl_color)
        sw.addWidget(more)
        sw.addStretch()
        g.addWidget(QLabel("Màu"), 0, 0)
        g.addLayout(sw, 0, 1, 1, 3)
        self.sl_op = QSlider(Qt.Horizontal)
        self.sl_op.setRange(10, 90)
        self.sl_op.setValue(45)
        self.lbl_op = QLabel("45%")
        self.sl_op.valueChanged.connect(lambda v: self.lbl_op.setText(f"{v}%"))
        g.addWidget(QLabel("Độ đậm"), 1, 0)
        g.addWidget(self.sl_op, 1, 1, 1, 2)
        g.addWidget(self.lbl_op, 1, 3)
        self.cb_hl_mode = QComboBox()
        self.cb_hl_mode.addItems(["Tự bắt chữ (không có chữ → tô vùng)", "Chỉ tô vùng tự do", "Chỉ tô chữ"])
        g.addWidget(QLabel("Kiểu"), 2, 0)
        g.addWidget(self.cb_hl_mode, 2, 1, 1, 3)
        self.ed_author = QLineEdit(os.environ.get("USERNAME", ""))
        g.addWidget(QLabel("Người tô"), 3, 0)
        g.addWidget(self.ed_author, 3, 1, 1, 3)
        lay.addWidget(grp)
        self.lbl_hl = QLabel("0 highlight")
        self.lbl_hl.setObjectName("summary")
        self.lbl_hl.setTextFormat(Qt.RichText)
        lay.addWidget(self.lbl_hl)
        self.hl_table = self._table(["", "Bản", "Trang", "Kiểu", "Text", "Ghi chú", "Người", "Ngày"], 4)
        for c, wd in enumerate([26, 50, 50, 50, 0, 140, 70, 110]):
            if wd:
                self.hl_table.setColumnWidth(c, wd)
        self.hl_table.itemClicked.connect(self._on_hl_click)
        self.hl_table.itemDoubleClicked.connect(lambda it: self._hl_comment(self._hl_by_row(it.row())))
        lay.addWidget(self.hl_table, 1)
        gg = QGridLayout()
        gg.setSpacing(6)
        btns = [("💬 Ghi chú", lambda: self._hl_comment(self._hl_selected())),
                ("🎨 Đổi màu", lambda: self._hl_recolor(self._hl_selected(), self.hl_color)),
                ("🗑 Xoá", self.hl_delete_selected), ("✕ Xoá hết", self.hl_clear),
                ("💾 Lưu vào PDF", self.hl_save_pdf), ("📊 Xuất Excel", self.hl_export),
                ("⤓ Lưu markup", self.hl_save_json), ("⤒ Mở markup", self.hl_load_json)]
        for i, (t, s) in enumerate(btns):
            b = QPushButton(t)
            if t.startswith("💾"):
                b.setObjectName("primary")
            b.clicked.connect(s)
            gg.addWidget(b, i // 4, i % 4)
        lay.addLayout(gg)
        return w

    def _build_status(self):
        sb = self.statusBar()
        self.lbl_coord = QLabel("X: –  Y: –")
        self.lbl_zoom = QLabel("Zoom: –")
        sb.addPermanentWidget(self.lbl_coord)
        sb.addPermanentWidget(self.lbl_zoom)

    def _build_shortcuts(self):
        sc = [("Ctrl+1", lambda: self.open_doc(1)), ("Ctrl+2", lambda: self.open_doc(2)),
              ("PgUp", lambda: self.set_page(self.page - 1)),
              ("PgDown", lambda: self.set_page(self.page + 1)),
              ("F", self.fit), ("H", lambda: self.set_tool("pan")), ("Esc", lambda: self.set_tool("pan")),
              ("C", lambda: self.set_tool("crop")), ("R", lambda: self.set_tool("read")),
              ("Z", lambda: self.set_tool("zone")), ("G", lambda: self.set_tool("highlight")),
              ("Ctrl+Z", self.hl_undo), ("Delete", self.hl_delete_selected),
              ("Ctrl+D", self.compare_current),
              ("+", lambda: self._zoom(1.25)), ("=", lambda: self._zoom(1.25)),
              ("-", lambda: self._zoom(0.8)),
              ("1", lambda: self.set_mode("v1")), ("2", lambda: self.set_mode("v2")),
              ("3", lambda: self.set_mode("side")), ("4", lambda: self.set_mode("overlay"))]
        for k, fn in sc:
            s = QShortcut(QKeySequence(k), self)
            s.setContext(Qt.WindowShortcut)
            s.activated.connect(fn)

    # ============================================================ helpers ==
    def offset(self) -> tuple[float, float]:
        return (self.sp_dx.value(), self.sp_dy.value())

    def both(self, page=None) -> bool:
        p = self.page if page is None else page
        d1, d2 = self.docs[1], self.docs[2]
        return bool(d1 and d2 and d1.has_page(p) and d2.has_page(p))

    def dpi_for(self, page: int) -> float:
        t = self.cb_dpi.currentText()
        if t != "Auto":
            return float(t)
        long = 0.0
        for d in self.docs.values():
            if d and d.has_page(page):
                r = d.page_rect(page)
                long = max(long, r.width, r.height)
        if not long:
            return 100.0
        return max(50.0, min(200.0, 5000 * 72 / long))

    def dpi(self) -> float:
        return self.dpi_for(self.page)

    def page_count(self) -> int:
        return max([d.page_count for d in self.docs.values() if d] or [0])

    def _to_view(self, r: QRectF, space, src) -> QRectF:
        dx, dy = self.offset()
        if src == 2 and space == 1:
            return r.translated(dx, dy)
        if src != 2 and space == 2:
            return r.translated(-dx, -dy)
        return QRectF(r)

    def _to_v1(self, r: QRectF, src) -> QRectF:
        if src == 2:
            dx, dy = self.offset()
            return r.translated(-dx, -dy)
        return QRectF(r)

    def msg(self, text, ms=5000):
        self.statusBar().showMessage(text, ms)

    def _invalidate(self, keep_diffs=False):
        self._pd_key = None
        self._pd = None
        self._td_cache.clear()
        if not keep_diffs:
            self.diff_items.clear()
            self.tree.clear()

    # ============================================================ actions ==
    def open_doc(self, ver: int, path: str | None = None):
        if not path:
            path, _ = QFileDialog.getOpenFileName(self, f"Mở bản vẽ Ver{ver}", self._last_dir, "PDF (*.pdf)")
        if not path:
            return
        try:
            with busy():
                doc = core.PdfDoc(path)
        except Exception as e:
            QMessageBox.critical(self, "Lỗi", f"Không mở được file:\n{e}")
            return
        pending = [h for h in self.highlights if h.ver == ver]
        if pending and self.docs[ver]:
            r = QMessageBox.question(self, "Highlight chưa lưu",
                                     f"PDF {ver} đang có {len(pending)} highlight. Mở file mới sẽ bỏ các highlight này.\n"
                                     "Tiếp tục? (Chọn No để quay lại và Lưu vào PDF trước)")
            if r != QMessageBox.Yes:
                doc.close()
                return
            self.highlights = [h for h in self.highlights if h.ver != ver]
            self._refresh_hl_table()
        if self.docs[ver]:
            self.docs[ver].close()
        self.docs[ver] = doc
        self._last_dir = os.path.dirname(path)
        self._invalidate()
        self.region = None
        self.region_td = None
        n = self.page_count()
        self.sp_page.blockSignals(True)
        self.sp_page.setRange(1, max(1, n))
        self.sp_page.blockSignals(False)
        self.lbl_pages.setText(f" / {n} ")
        if self.page >= n:
            self.page = 0
        if self.docs[1] and self.docs[2] and self.mode in ("v1", "v2"):
            self.set_mode("side")
        elif ver == 2 and not self.docs[1]:
            self.set_mode("v2")
        else:
            self.refresh_view()
        self.msg(f"Đã mở Ver{ver}: {doc.name} ({doc.page_count} trang)")

    def set_page(self, p: int):
        n = self.page_count()
        if n == 0:
            return
        p = max(0, min(n - 1, p))
        if p == self.page and self.view_src and self.viewA.pix_item is not None:
            return
        self.page = p
        self.sp_page.blockSignals(True)
        self.sp_page.setValue(p + 1)
        self.sp_page.blockSignals(False)
        self.refresh_view()

    def set_mode(self, mode: str):
        self.mode = mode
        self.mode_btns[mode].setChecked(True)
        self.viewB.setVisible(mode == "side")
        if mode == "side":
            half = max(200, self.splitter.width() // 2)
            self.splitter.setSizes([half, half])
            QApplication.processEvents()
        self.refresh_view()
        if mode == "side":
            QTimer.singleShot(0, lambda: self._sync(self.viewA, self.viewB))

    def set_tool(self, tool: str):
        self.tool = tool
        self.tool_btns[tool].setChecked(True)
        for v in (self.viewA, self.viewB):
            v.set_tool(tool, self.hl_color if tool == "highlight" else TOOL_COL[tool])
        if tool == "highlight":
            self.tabs.setCurrentIndex(4)
        self.msg(TOOL_HINT[tool], 8000)

    def fit(self):
        self.viewA.fit()
        if self.viewB.isVisible():
            self._sync(self.viewA, self.viewB)

    def _zoom(self, f):
        self.viewA.zoom_by(f)

    # ========================================================== rendering ==
    def refresh_view(self):
        dpi = self.dpi()
        srcs = {"v1": [(self.viewA, 1)], "v2": [(self.viewA, 2)],
                "side": [(self.viewA, 1), (self.viewB, 2)],
                "overlay": [(self.viewA, "ov")]}[self.mode]
        self.view_src = {}
        with busy():
            for view, src in srcs:
                self.view_src[view] = src
                self._show(view, src, dpi)
        self.redraw_overlays()
        self._hires_timer.start()

    def _show(self, view: PdfView, src, dpi: float):
        z = dpi / 72.0
        p = self.page
        if src == "ov":
            view.title = "OVERLAY"
            view.accent = ORANGE
            view.subtitle = f"Trang {p + 1}  ·  Đỏ = chỉ Ver1  ·  Xanh = chỉ Ver2"
            pd = self.get_pdiff()
            if pd is None:
                view.placeholder = "Overlay cần mở cả Ver1 và Ver2 (cùng có trang này)"
                view.set_image(None, 1, QRectF())
                return
            arr = pd.overlay_rgb()
            h, w = arr.shape[:2]
            view.set_image(np_to_qimage(arr), z, QRectF(0, 0, w / z, h / z))
            return
        doc = self.docs[src]
        view.title = f"VER {src}"
        view.accent = BLUE if src == 1 else GREEN
        if doc is None:
            view.subtitle = ""
            view.placeholder = f"Kéo thả file PDF vào đây hoặc bấm 📂 PDF {src}"
            view.set_image(None, 1, QRectF())
            return
        view.subtitle = f"{doc.name}  ·  Trang {p + 1}/{doc.page_count}"
        if not doc.has_page(p):
            view.placeholder = f"Ver{src} không có trang {p + 1}"
            view.set_image(None, 1, QRectF())
            return
        arr = doc.render(p, dpi)
        r = doc.page_rect(p)
        view.set_image(np_to_qimage(arr), z, QRectF(0, 0, r.width, r.height))

    def get_pdiff(self):
        if not self.both():
            return None
        d1, d2 = self.docs[1], self.docs[2]
        p, dpi, off = self.page, self.dpi(), self.offset()
        key = (p, round(dpi, 2), self.sp_thr.value(), self.sp_tol.value(), off)
        if key != self._pd_key:
            z = dpi / 72.0
            with busy():
                self._pd = core.pixel_diff(d1.render(p, dpi, True), d2.render(p, dpi, True), dpi,
                                           self.sp_thr.value(), self.sp_tol.value(),
                                           (int(round(-off[0] * z)), int(round(-off[1] * z))))
            self._pd_key = key
        return self._pd

    def _do_hires(self):
        """Render a sharp tile of the visible area when zoomed beyond base DPI."""
        p = self.page
        for view, src in self.view_src.items():
            if not view.isVisible() or view.pix_item is None:
                continue
            view.clear_layer("hires")
            base = self.dpi()
            need = view.zoom_factor() * 72 * view.devicePixelRatioF()
            if need <= base * 1.15:
                continue
            need = min(need, 1600)
            vis = view.mapToScene(view.viewport().rect()).boundingRect()
            if src == "ov":
                d1, d2 = self.docs[1], self.docs[2]
                if not self.both():
                    continue
                pr = QRectF(0, 0, d1.page_rect(p).width, d1.page_rect(p).height)
            else:
                d = self.docs[src]
                pr = QRectF(0, 0, d.page_rect(p).width, d.page_rect(p).height)
            vis = vis.intersected(pr)
            if vis.isEmpty() or vis.width() * vis.height() * (need / 72) ** 2 > 45e6:
                continue
            clip = frect(vis)
            try:
                if src == "ov":
                    dx, dy = self.offset()
                    g1 = d1.render(p, need, True, clip)
                    g2 = d2.render(p, need, True, fitz.Rect(clip.x0 + dx, clip.y0 + dy, clip.x1 + dx, clip.y1 + dy))
                    arr = core.pixel_diff(g1, g2, need, self.sp_thr.value(), self.sp_tol.value()).overlay_rgb()
                else:
                    arr = self.docs[src].render(p, need, False, clip)
            except Exception:
                continue
            view.add_pixmap("hires", np_to_qimage(arr), need / 72.0, vis.topLeft(), z=1)

    # ============================================================ overlays ==
    def redraw_overlays(self):
        p = self.page
        for view, src in self.view_src.items():
            for tag in ("diff", "zones", "query", "region", "focus"):
                view.clear_layer(tag)
            if view.pix_item is None:
                continue
            if self.chk_show.isChecked():
                for it in self.diff_items.get(p, []):
                    if src == 1 and it.kind in ("added", "geo_added"):
                        continue
                    if src == 2 and it.kind in ("removed", "geo_removed"):
                        continue
                    geo = it.kind.startswith("geo")
                    r = self._to_view(qrect(it.rect), 1, src)
                    if geo:
                        view.add_rect("diff", r, KIND_COL[it.kind], 14, 1.6, dash=True, z=8)
                    else:
                        view.add_rect("diff", r.adjusted(-1, -1, 1, 1), KIND_COL[it.kind], 70, 1.2, z=9)
            for zn in self.zones:
                if zn.page != p:
                    continue
                lab = zn.name + (f" · {zn.status}" if zn.status else "")
                view.add_rect("zones", self._to_view(qrect(zn.rect), 1, src), ZONE_COL.get(zn.status, PURPLE),
                              22, 2.2, label=lab, z=12)
            for h in self.query_hits:
                if h.page != p or (src in (1, 2) and h.ver != src):
                    continue
                col = (RED if h.ver == 1 else GREEN) if h.changed else self.query_color
                view.add_rect("query", self._to_view(qrect(h.rect), h.ver, src).adjusted(-1.5, -1.5, 1.5, 1.5),
                              col, 90, 2, z=14)
            if self.region and self.region[0] == p:
                view.add_rect("region", self._to_view(self.region[1], 1, src), YELLOW, 8, 2,
                              label="Vùng đọc data", dash=True, z=11)
                td = self.region_td
                if td:
                    dx, dy = self.offset()
                    if src in (1, "ov"):
                        for w in td.removed:
                            view.add_rect("region", self._to_view(qrect(w[:4]), 1, src), RED, 90, 1, z=13)
                    if src in (2, "ov"):
                        for w in td.added:
                            view.add_rect("region", self._to_view(qrect(w[:4]), 2, src), GREEN, 90, 1, z=13)
                    for a, b in td.changed:
                        if src in (1, "ov"):
                            view.add_rect("region", self._to_view(qrect(a[:4]), 1, src), ORANGE, 90, 1, z=13)
                        if src == 2:
                            view.add_rect("region", qrect(b[:4]), ORANGE, 90, 1, z=13)
            for h in self.highlights:
                if h.page != p or (src in (1, 2) and h.ver != src):
                    continue
                col = QColor(h.color)
                for r in h.rects:
                    view.add_highlight("region", self._to_view(qrect(r), h.ver, src), col, h.opacity)
                if h.id == self.hl_sel:
                    view.add_rect("region", self._to_view(qrect(h.bbox()), h.ver, src).adjusted(-2, -2, 2, 2),
                                  col.darker(160), 0, 1.5, dash=True, z=7)
                if h.comment:
                    b = self._to_view(qrect(h.bbox()), h.ver, src)
                    view.add_rect("region", QRectF(b.right() - 1, b.top() - 1, 2, 2), col.darker(130), 0, 1,
                                  label="💬", z=7)

    def _focus(self, r: QRectF, space=1):
        for view, src in self.view_src.items():
            view.clear_layer("focus")
            view.add_rect("focus", self._to_view(r, space, src).adjusted(-3, -3, 3, 3), CYAN, 0, 3, z=20)

    def _goto(self, page: int, r: QRectF, space=1, margin=2.5):
        if page != self.page:
            self.set_page(page)
        src = self.view_src.get(self.viewA, 1)
        self.viewA.zoom_to(self._to_view(r, space, src), margin)
        if self.viewB.isVisible():
            self._sync(self.viewA, self.viewB)
        self._focus(r, space)

    # ============================================================== events ==
    def on_view_changed(self, view: PdfView):
        if self.mode == "side":
            other = self.viewB if view is self.viewA else self.viewA
            self._sync(view, other)
        self.lbl_zoom.setText(f"Zoom: {view.zoom_factor() * 100:.0f}%")
        self._hires_timer.start()

    def _sync(self, src: PdfView, dst: PdfView):
        if self._syncing or not dst.isVisible() or src.pix_item is None:
            return
        self._syncing = True
        try:
            dst.setTransform(src.transform())
            c = src.mapToScene(src.viewport().rect().center())
            dx, dy = self.offset()
            c = c + QPointF(dx, dy) if src is self.viewA else c - QPointF(dx, dy)
            dst.centerOn(c)
        finally:
            self._syncing = False

    def on_mouse(self, p: QPointF):
        self.lbl_coord.setText(f"X: {p.x():.1f}  Y: {p.y():.1f} pt   ({p.x() * 25.4 / 72:.1f}, "
                               f"{p.y() * 25.4 / 72:.1f} mm)")

    def on_rect(self, view: PdfView, r: QRectF):
        src = self.view_src.get(view, 1)
        if self.tool == "highlight":
            self.hl_add_from_rect(view, r)
            return
        r1 = self._to_v1(r, src)
        if self.tool == "crop":
            self.do_crop(r1)
        elif self.tool == "read":
            self.region = (self.page, r1)
            self.run_read()
            self.tabs.setCurrentIndex(1)
        elif self.tool == "zone":
            self.add_zone(r1)

    def _settings_applied(self):
        self._invalidate(keep_diffs=True)
        self.refresh_view()
        if self.region:
            self.run_read()

    def dragEnterEvent(self, e):
        if e.mimeData().hasUrls() and any(u.toLocalFile().lower().endswith(".pdf") for u in e.mimeData().urls()):
            e.acceptProposedAction()

    def dropEvent(self, e):
        paths = [u.toLocalFile() for u in e.mimeData().urls() if u.toLocalFile().lower().endswith(".pdf")]
        if len(paths) >= 2:
            paths.sort()
            self.open_doc(1, paths[0])
            self.open_doc(2, paths[1])
            return
        if not paths:
            return
        if not self.docs[1]:
            self.open_doc(1, paths[0])
        elif not self.docs[2]:
            self.open_doc(2, paths[0])
        else:
            mb = QMessageBox(self)
            mb.setWindowTitle("Mở file")
            mb.setText(f"Mở '{os.path.basename(paths[0])}' làm:")
            b1 = mb.addButton("Ver1", QMessageBox.AcceptRole)
            b2 = mb.addButton("Ver2", QMessageBox.AcceptRole)
            mb.addButton("Huỷ", QMessageBox.RejectRole)
            mb.exec()
            if mb.clickedButton() is b1:
                self.open_doc(1, paths[0])
            elif mb.clickedButton() is b2:
                self.open_doc(2, paths[0])

    # ============================================================ compare ==
    def _need_both(self) -> bool:
        if not (self.docs[1] and self.docs[2]):
            QMessageBox.information(self, "So sánh", "Cần mở cả Ver1 và Ver2.")
            return False
        return True

    def compare_current(self):
        if not self._need_both():
            return
        if not self.both():
            QMessageBox.information(self, "So sánh", f"Một trong hai bản không có trang {self.page + 1}.")
            return
        with busy():
            pd = self.get_pdiff() if self.chk_pixel.isChecked() else None
            items, _ = core.compare_page(self.docs[1], self.docs[2], self.page, self.dpi(),
                                         self.sp_thr.value(), self.sp_tol.value(), self.offset(),
                                         self.chk_pixel.isChecked(), self.chk_text.isChecked(), pdiff=pd)
        self.diff_items[self.page] = items
        self._populate_tree()
        self.redraw_overlays()
        self.tabs.setCurrentIndex(0)
        self.msg(f"Trang {self.page + 1}: {len(items)} khác biệt")

    def compare_all(self):
        if not self._need_both():
            return
        n = min(self.docs[1].page_count, self.docs[2].page_count)
        dlg = QProgressDialog("Đang so sánh…", "Huỷ", 0, n, self)
        dlg.setWindowTitle("So sánh tất cả trang")
        dlg.setWindowModality(Qt.WindowModal)
        dlg.setMinimumDuration(0)
        for p in range(n):
            dlg.setValue(p)
            dlg.setLabelText(f"Đang so sánh trang {p + 1}/{n}…")
            QApplication.processEvents()
            if dlg.wasCanceled():
                break
            items, _ = core.compare_page(self.docs[1], self.docs[2], p, min(self.dpi_for(p), 100),
                                         self.sp_thr.value(), self.sp_tol.value(), self.offset(),
                                         self.chk_pixel.isChecked(), self.chk_text.isChecked())
            self.diff_items[p] = items
        dlg.setValue(n)
        self._populate_tree()
        self.redraw_overlays()
        self.tabs.setCurrentIndex(0)

    def _populate_tree(self):
        self.tree.clear()
        tot = Counter()
        bold = QFont("Segoe UI", 10, QFont.Bold)
        for p in sorted(self.diff_items):
            items = self.diff_items[p]
            top = QTreeWidgetItem([f"Trang {p + 1}", f"{len(items)} khác biệt", ""])
            top.setData(0, Qt.UserRole, (p, -1))
            top.setFont(0, bold)
            top.setForeground(1, QBrush(ORANGE if items else GREEN))
            for i, it in enumerate(items):
                c = QTreeWidgetItem([it.label, it.text, f"{it.rect[0]:.0f}, {it.rect[1]:.0f}"])
                c.setData(0, Qt.UserRole, (p, i))
                c.setForeground(0, QBrush(KIND_COL[it.kind]))
                top.addChild(c)
                tot[it.kind] += 1
            self.tree.addTopLevelItem(top)
            top.setExpanded(p == self.page or len(self.diff_items) == 1)
        geo = tot["geo_removed"] + tot["geo_added"] + tot["geo_changed"]
        self.lbl_cmp.setText(
            f"<b>{len(self.diff_items)}</b> trang đã so sánh &nbsp;·&nbsp; "
            f"<span style='color:{ORANGE.name()}'>Nét vẽ: <b>{geo}</b> vùng</span><br>"
            f"<span style='color:{RED.name()}'>Text xoá: <b>{tot['removed']}</b></span> &nbsp; "
            f"<span style='color:{GREEN.name()}'>Text thêm: <b>{tot['added']}</b></span> &nbsp; "
            f"<span style='color:{ORANGE.name()}'>Text sửa: <b>{tot['changed']}</b></span>")

    def on_tree_click(self, item: QTreeWidgetItem, _col=0):
        data = item.data(0, Qt.UserRole)
        if not data:
            return
        p, i = data
        if i < 0:
            if p != self.page:
                self.set_page(p)
            return
        it = self.diff_items[p][i]
        self._goto(p, qrect(it.rect), 1, 3.0)

    def export_diffs(self):
        rows = []
        for p in sorted(self.diff_items):
            for it in self.diff_items[p]:
                rows.append([p + 1, it.label, it.text] + [round(v, 1) for v in it.rect])
        if not rows:
            QMessageBox.information(self, "Xuất", "Chưa có kết quả so sánh.")
            return
        self.write_table("Khac biet", ["Trang", "Loại", "Nội dung", "x0", "y0", "x1", "y1"], rows,
                         "diff_report.xlsx", status_col=1)

    # ========================================================== read data ==
    def run_read(self):
        if not self.region:
            return
        p, r1 = self.region
        d1, d2 = self.docs[1], self.docs[2]
        dx, dy = self.offset()
        f1 = frect(r1)
        f2 = fitz.Rect(f1.x0 + dx, f1.y0 + dy, f1.x1 + dx, f1.y1 + dy)
        has1 = bool(d1 and d1.has_page(p))
        has2 = bool(d2 and d2.has_page(p))
        lines1 = d1.lines_in(p, f1) if has1 else None
        lines2 = d2.lines_in(p, f2) if has2 else None
        if has1 and has2:
            rows = core.line_diff(lines1, lines2)
            self.region_td = core.text_diff(d1.words_in(p, f1), d2.words_in(p, f2), self.offset())
        else:
            rows = [("equal", l, "") for l in (lines1 or [])] + [("equal", "", l) for l in (lines2 or [])]
            self.region_td = None
        self.read_table.setRowCount(0)
        self.read_table.setColumnHidden(0, not (has1 and has2))
        self.read_table.setColumnHidden(1, not has1)
        self.read_table.setColumnHidden(2, not has2)
        for st, a, b in rows:
            r = self.read_table.rowCount()
            self.read_table.insertRow(r)
            for c, txt in enumerate([ROW_SYM[st] if (has1 and has2) else "", a, b]):
                it = QTableWidgetItem(txt)
                it.setToolTip(txt)
                if ROW_BG[st] is not None:
                    it.setBackground(ROW_BG[st])
                self.read_table.setItem(r, c, it)
        # preview
        long = max(f1.width, f1.height)
        pdpi = max(72.0, min(600.0, 900 * 72 / max(long, 1)))
        info_px = ""
        try:
            if has1 and has2:
                pd = core.pixel_diff(d1.render(p, pdpi, True, f1), d2.render(p, pdpi, True, f2), pdpi,
                                     self.sp_thr.value(), self.sp_tol.value())
                arr = pd.overlay_rgb()
                info_px = (f" &nbsp;·&nbsp; Nét vẽ: <span style='color:{RED.name()}'>−{pd.n_removed}</span> / "
                           f"<span style='color:{GREEN.name()}'>+{pd.n_added}</span> px")
            else:
                d, f = (d1, f1) if has1 else (d2, f2)
                arr = d.render(p, pdpi, False, f) if d else None
            if arr is not None:
                pm = QPixmap.fromImage(np_to_qimage(arr))
                self.read_preview.setPixmap(pm.scaled(self.read_preview.width() - 8, self.read_preview.height() - 8,
                                                      Qt.KeepAspectRatio, Qt.SmoothTransformation))
        except Exception as e:
            self.read_preview.setText(f"Lỗi preview: {e}")
        cnt = Counter(st for st, _, _ in rows)
        self.lbl_read.setText(
            f"Trang <b>{p + 1}</b> &nbsp;·&nbsp; vùng {f1.width:.0f}×{f1.height:.0f} pt "
            f"({f1.width * 25.4 / 72:.0f}×{f1.height * 25.4 / 72:.0f} mm)<br>"
            f"Ver1: <b>{len(lines1) if lines1 is not None else '–'}</b> dòng &nbsp; "
            f"Ver2: <b>{len(lines2) if lines2 is not None else '–'}</b> dòng &nbsp;·&nbsp; "
            f"<span style='color:{ORANGE.name()}'>≠ {cnt['changed']}</span> "
            f"<span style='color:{RED.name()}'>− {cnt['removed']}</span> "
            f"<span style='color:{GREEN.name()}'>+ {cnt['added']}</span>{info_px}")
        self.redraw_overlays()

    def _read_rows(self):
        rows = []
        for r in range(self.read_table.rowCount()):
            rows.append([self.read_table.item(r, c).text() if self.read_table.item(r, c) else ""
                         for c in range(3)])
        return rows

    def read_copy(self):
        rows = self._read_rows()
        QApplication.clipboard().setText("\n".join("\t".join(r) for r in [["", "Ver1", "Ver2"]] + rows))
        self.msg(f"Đã copy {len(rows)} dòng vào clipboard")

    def read_export(self):
        rows = self._read_rows()
        if not rows:
            return
        self.write_table("Doc data", ["Trạng thái", "Ver1", "Ver2"], rows, "region_data.xlsx")

    # ============================================================== crop ==
    def do_crop(self, r1: QRectF):
        d1, d2, p = self.docs[1], self.docs[2], self.page
        has1 = bool(d1 and d1.has_page(p))
        has2 = bool(d2 and d2.has_page(p))
        if not (has1 or has2):
            return
        dlg = CropDialog(self, has1, has2)
        if not dlg.exec():
            return
        self._export_crops([(p, r1, "")], dlg.chk1.isChecked(), dlg.chk2.isChecked(), dlg.chkd.isChecked(),
                           dlg.rb_pdf.isChecked(), dlg.sp_dpi.value())

    def _export_crops(self, regions, use1, use2, use_diff, as_pdf, dpi):
        d1, d2 = self.docs[1], self.docs[2]
        dx, dy = self.offset()
        base = os.path.join(self._last_dir, f"crop_p{regions[0][0] + 1}")
        if as_pdf:
            path, _ = QFileDialog.getSaveFileName(self, "Lưu crop PDF", base + ".pdf", "PDF (*.pdf)")
        else:
            path, _ = QFileDialog.getSaveFileName(self, "Lưu crop PNG (tên gốc)", base + ".png", "PNG (*.png)")
        if not path:
            return
        stem = os.path.splitext(path)[0]
        out = []
        try:
            with busy():
                items = []
                for p, r1, name in regions:
                    f1 = frect(r1)
                    f2 = fitz.Rect(f1.x0 + dx, f1.y0 + dy, f1.x1 + dx, f1.y1 + dy)
                    sfx = f"_{name}" if name else ""
                    if use1 and d1 and d1.has_page(p):
                        items.append((d1, p, f1, f"{sfx}_v1"))
                    if use2 and d2 and d2.has_page(p):
                        items.append((d2, p, f2, f"{sfx}_v2"))
                    if use_diff and self.both(p):
                        arr = core.pixel_diff(d1.render(p, dpi, True, f1), d2.render(p, dpi, True, f2), dpi,
                                              self.sp_thr.value(), self.sp_tol.value()).overlay_rgb()
                        fp = f"{stem}{sfx}_diff.png"
                        core.save_rgb_png(arr, fp)
                        out.append(fp)
                if as_pdf and items:
                    core.export_crop_pdf(path, [(d, p, f) for d, p, f, _ in items])
                    out.insert(0, path)
                elif not as_pdf:
                    for d, p, f, sfx in items:
                        fp = f"{stem}{sfx}.png"
                        d.page(p).get_pixmap(dpi=dpi, clip=f, alpha=False).save(fp)
                        out.append(fp)
        except Exception as e:
            QMessageBox.critical(self, "Crop", f"Lỗi: {e}")
            return
        if not out:
            return
        self._last_dir = os.path.dirname(path)
        r = QMessageBox.question(self, "Crop xong", "Đã lưu:\n" + "\n".join(out) + "\n\nMở file?")
        if r == QMessageBox.Yes:
            os.startfile(out[0])

    # ============================================================= zones ==
    def add_zone(self, r1: QRectF):
        p = self.page
        d = self.docs[2] if self.docs[2] and self.docs[2].has_page(p) else self.docs[1]
        preview = ""
        if d and d.has_page(p):
            rr = r1 if d is self.docs[1] else self._to_view(r1, 1, 2)
            preview = "\n".join(d.lines_in(p, frect(rr)))
        dlg = ZoneDialog(self, f"Z{len(self.zones) + 1:02d}", preview=preview,
                         allow_compare=bool(self.docs[1] and self.docs[2]))
        if not dlg.exec():
            return
        name, rule, exp = dlg.values()
        z = core.Zone(name, p, ttuple(r1), rule, exp)
        self._check(z)
        self.zones.append(z)
        self._refresh_zone_table()
        self.redraw_overlays()
        self.tabs.setCurrentIndex(2)

    def _check(self, z: core.Zone):
        core.check_zone(z, self.docs[1], self.docs[2], self.offset(), 200,
                        self.sp_thr.value(), self.sp_tol.value())

    def _templates(self) -> list[core.Zone]:
        """One zone per name (first occurrence) = the zone definitions."""
        seen, out = set(), []
        for z in self.zones:
            if z.name not in seen:
                seen.add(z.name)
                out.append(z)
        return out

    def _expand_all_pages(self) -> int:
        n = self.page_count()
        exist = {(z.name, z.page) for z in self.zones}
        added = 0
        for t in self._templates():
            for p in range(n):
                if (t.name, p) not in exist:
                    self.zones.append(core.Zone(t.name, p, t.rect, t.rule, t.expected))
                    exist.add((t.name, p))
                    added += 1
        order = {t.name: i for i, t in enumerate(self._templates())}
        self.zones.sort(key=lambda z: (z.page, order.get(z.name, 0)))
        return added

    def run_zone_check(self):
        if not self.zones:
            QMessageBox.information(self, "Vùng check", "Chưa có vùng nào. Dùng công cụ ▣ Vùng check (Z) để tạo.")
            return
        if not (self.docs[1] or self.docs[2]):
            QMessageBox.information(self, "Vùng check", "Chưa mở file PDF nào.")
            return
        if self.chk_allpages.isChecked():
            self._expand_all_pages()
        with busy():
            for z in self.zones:
                self._check(z)
        self._refresh_zone_table()
        self.redraw_overlays()
        cnt = Counter(z.status for z in self.zones)
        self.msg(f"Đã kiểm tra {len(self.zones)} vùng trên {len({z.page for z in self.zones})} trang — "
                 + ", ".join(f"{k}: {v}" for k, v in cnt.items()), 10000)

    def _refresh_zone_table(self):
        t = self.zone_table
        t.setRowCount(0)
        cnt = Counter()
        for z in self.zones:
            r = t.rowCount()
            t.insertRow(r)
            vals = [z.name, str(z.page + 1), core.RULES.get(z.rule, z.rule), z.expected, z.status or "—",
                    z.detail, z.text_v1.replace("\n", " ⏎ "), z.text_v2.replace("\n", " ⏎ ")]
            for c, v in enumerate(vals):
                it = QTableWidgetItem(v)
                it.setToolTip(v if c < 6 else (z.text_v1 if c == 6 else z.text_v2))
                if c == 4 and z.status:
                    col = QColor(ZONE_COL.get(z.status, GREY))
                    it.setForeground(QBrush(col))
                    bg = QColor(col)
                    bg.setAlpha(45)
                    it.setBackground(bg)
                    f = it.font()
                    f.setBold(True)
                    it.setFont(f)
                t.setItem(r, c, it)
            cnt[z.status or "—"] += 1
        parts = [f"<b>{len(self.zones)}</b> vùng"]
        for st in ("OK", "CHANGED", "FAIL", "DATA", "N/A"):
            if cnt[st]:
                parts.append(f"<span style='color:{ZONE_COL[st].name()}'>{st}: <b>{cnt[st]}</b></span>")
        self.lbl_zone.setText(" &nbsp;·&nbsp; ".join(parts))
        t.setColumnHidden(6, self.docs[1] is None and self.docs[2] is not None)
        t.setColumnHidden(7, self.docs[2] is None and self.docs[1] is not None)

    def _sel_zone_rows(self):
        return sorted({i.row() for i in self.zone_table.selectedIndexes()})

    def on_zone_click(self, item):
        z = self.zones[item.row()]
        self._goto(z.page, qrect(z.rect), 1, 1.6)

    def edit_zone(self):
        rows = self._sel_zone_rows()
        if not rows:
            return
        z = self.zones[rows[0]]
        dlg = ZoneDialog(self, z.name, z.rule, z.expected, z.text_v2 or z.text_v1,
                         allow_compare=bool(self.docs[1] and self.docs[2]))
        if dlg.exec():
            old = z.name
            name, rule, exp = dlg.values()
            family = [x for x in self.zones if x.name == old] if self.chk_allpages.isChecked() else [z]
            for x in family:
                x.name, x.rule, x.expected = name, rule, exp
                self._check(x)
            self._refresh_zone_table()
            self.redraw_overlays()

    def delete_zones(self):
        rows = self._sel_zone_rows()
        if self.chk_allpages.isChecked():
            names = {self.zones[r].name for r in rows}
            self.zones = [z for z in self.zones if z.name not in names]
        else:
            for r in reversed(rows):
                del self.zones[r]
        self._refresh_zone_table()
        self.redraw_overlays()

    def save_template(self):
        if not self.zones:
            return
        path, _ = QFileDialog.getSaveFileName(self, "Lưu template vùng check",
                                              os.path.join(self._last_dir, "qa_template.json"), "JSON (*.json)")
        if not path:
            return
        src = self._templates() if self.chk_allpages.isChecked() else self.zones
        data = {"app": "PDF QA Viewer", "version": 1, "offset": list(self.offset()),
                "all_pages": self.chk_allpages.isChecked(),
                "zones": [z.to_json() for z in src]}
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        self.msg(f"Đã lưu template: {path}")

    def load_template(self):
        path, _ = QFileDialog.getOpenFileName(self, "Mở template vùng check", self._last_dir, "JSON (*.json)")
        if not path:
            return
        try:
            with open(path, encoding="utf-8") as f:
                data = json.load(f)
            zones = [core.Zone.from_json(d) for d in data.get("zones", [])]
        except Exception as e:
            QMessageBox.critical(self, "Template", f"File không hợp lệ: {e}")
            return
        if self.zones:
            r = QMessageBox.question(self, "Template", "Thay thế các vùng hiện có?\n(No = thêm vào)",
                                     QMessageBox.Yes | QMessageBox.No | QMessageBox.Cancel)
            if r == QMessageBox.Cancel:
                return
            if r == QMessageBox.Yes:
                self.zones = []
        self.zones.extend(zones)
        self.run_zone_check() if (self.docs[1] or self.docs[2]) else self._refresh_zone_table()
        self.redraw_overlays()
        self.tabs.setCurrentIndex(2)

    def export_zones(self):
        if not self.zones:
            return
        rows = [[z.name, z.page + 1, core.RULES.get(z.rule, z.rule), z.expected, z.status, z.detail,
                 z.text_v1, z.text_v2, z.geo_changed_px] + [round(v, 1) for v in z.rect] for z in self.zones]
        self.write_table("Vung check", ["Tên", "Trang", "Quy tắc", "Giá trị", "Kết quả", "Chi tiết",
                                        "Text Ver1", "Text Ver2", "Px hình khác", "x0", "y0", "x1", "y1"],
                         rows, "zone_check.xlsx", status_col=4)

    def replicate_zones(self):
        self.chk_allpages.setChecked(True)
        self.run_zone_check()

    # ------------------------------------------------------- batch files --
    def batch_check_files(self):
        temps = self._templates()
        if not temps:
            QMessageBox.information(self, "Kiểm tra nhiều file",
                                    "Chưa có vùng nào. Mở 1 bản vẽ mẫu, dùng ▣ Vùng check (Z) để tạo vùng "
                                    "(hoặc 📂 Mở template) rồi chạy lại.")
            return
        paths, _ = QFileDialog.getOpenFileNames(self, "Chọn các file PDF cần kiểm tra (Ctrl+A để chọn tất cả)",
                                                self._last_dir, "PDF (*.pdf)")
        if not paths:
            return
        self._last_dir = os.path.dirname(paths[0])
        dlg = QProgressDialog("Đang kiểm tra…", "Huỷ", 0, len(paths), self)
        dlg.setWindowTitle("Kiểm tra nhiều file")
        dlg.setWindowModality(Qt.WindowModal)
        dlg.setMinimumDuration(0)
        results = []   # (path, Zone)
        for i, path in enumerate(paths):
            dlg.setValue(i)
            dlg.setLabelText(f"{i + 1}/{len(paths)}: {os.path.basename(path)}")
            QApplication.processEvents()
            if dlg.wasCanceled():
                break
            try:
                doc = core.PdfDoc(path)
            except Exception as e:
                z = core.Zone("(file)", 0, (0, 0, 0, 0), "extract")
                z.status, z.detail = "FAIL", f"Không mở được: {e}"
                results.append((path, z))
                continue
            for p in range(doc.page_count):
                pr = doc.page_rect(p)
                for t in temps:
                    rule = "extract" if t.rule == "compare" else t.rule
                    z = core.Zone(t.name, p, t.rect, rule, t.expected)
                    if not fitz.Rect(t.rect).intersects(pr):
                        z.status, z.detail = "N/A", "vùng nằm ngoài trang (khác khổ giấy?)"
                    else:
                        core.check_zone(z, doc, None)
                    results.append((path, z))
            doc.close()
        dlg.setValue(len(paths))
        if results:
            BatchResultDialog(self, results, [t.name for t in temps]).exec()

    def export_zone_pivot(self):
        """Data table: one row per page, one column per zone name (value = text of the zone)."""
        if not self.zones:
            return
        names = list(dict.fromkeys(z.name for z in self.zones))
        pages = sorted({z.page for z in self.zones})
        cell = {(z.page, z.name): z for z in self.zones}
        rows = []
        for p in pages:
            row = [p + 1]
            st = []
            for nm in names:
                z = cell.get((p, nm))
                row.append(((z.text_v2 or z.text_v1).replace("\n", " ")) if z else "")
                if z and z.status in ("FAIL", "CHANGED"):
                    st.append(f"{nm}:{z.status}")
            row.append("; ".join(st) or "OK")
            rows.append(row)
        self.write_table("Bang data", ["Trang"] + names + ["Kiểm tra"], rows, "zone_data.xlsx",
                         status_col=len(names) + 1)

    def crop_zones(self):
        rows = self._sel_zone_rows() or list(range(len(self.zones)))
        if not rows:
            return
        d1, d2 = self.docs[1], self.docs[2]
        dlg = CropDialog(self, bool(d1), bool(d2))
        if not dlg.exec():
            return
        regs = [(self.zones[i].page, qrect(self.zones[i].rect), self.zones[i].name) for i in rows]
        self._export_crops(regs, dlg.chk1.isChecked(), dlg.chk2.isChecked(), dlg.chkd.isChecked(),
                           dlg.rb_pdf.isChecked(), dlg.sp_dpi.value())

    # ============================================================= query ==
    def _upd_qcol(self):
        self.btn_qcol.setStyleSheet(f"QPushButton {{ color: {self.query_color.name()}; font-weight:600; }}")

    def pick_query_color(self):
        c = QColorDialog.getColor(self.query_color, self, "Màu highlight")
        if c.isValid():
            self.query_color = c
            self._upd_qcol()
            self.redraw_overlays()

    def _changed_keys(self, p: int):
        key = (p, self.offset())
        if key not in self._td_cache:
            td = core.text_diff(self.docs[1].words(p), self.docs[2].words(p), self.offset())
            k1 = {tuple(w[:4]) for w in td.removed} | {tuple(a[:4]) for a, _ in td.changed}
            k2 = {tuple(w[:4]) for w in td.added} | {tuple(b[:4]) for _, b in td.changed}
            self._td_cache[key] = (k1, k2)
        return self._td_cache[key]

    def run_query(self):
        pat = self.q_edit.text().strip()
        if not pat:
            return
        scope = self.cb_scope.currentIndex()
        vers = {0: [1, 2], 1: [1], 2: [2]}[self.cb_ver.currentIndex()]
        pages = [self.page] if scope == 0 else list(range(self.page_count()))
        rects = None
        if scope == 2:
            rects = {}
            for z in self.zones:
                rects.setdefault(z.page, []).append(z.rect)
            pages = sorted(rects)
        elif scope == 3:
            if not self.region:
                QMessageBox.information(self, "Query", "Chưa có vùng Đọc data. Dùng công cụ 🔎 (R) để chọn.")
                return
            pages = [self.region[0]]
            rects = {self.region[0]: [ttuple(self.region[1])]}
        dx, dy = self.offset()
        hits = []
        with busy():
            for v in vers:
                d = self.docs[v]
                if not d:
                    continue
                rr = rects
                if rects is not None and v == 2:
                    rr = {p: [(r[0] + dx, r[1] + dy, r[2] + dx, r[3] + dy) for r in lst] for p, lst in rects.items()}
                hits += core.query_words(d, v, pages, pat, self.chk_regex.isChecked(),
                                         self.chk_case.isChecked(), rr)
            for h in hits:
                if self.both(h.page):
                    k1, k2 = self._changed_keys(h.page)
                    h.changed = tuple(h.rect) in (k1 if h.ver == 1 else k2)
        if self.chk_changed.isChecked():
            hits = [h for h in hits if h.changed]
        hits.sort(key=lambda h: (h.page, h.ver, h.rect[1], h.rect[0]))
        self.query_hits = hits
        t = self.q_table
        t.setRowCount(0)
        for h in hits:
            r = t.rowCount()
            t.insertRow(r)
            st = ("Xoá ở Ver2" if h.ver == 1 else "Mới ở Ver2") if h.changed else "Không đổi"
            vals = [f"V{h.ver}", str(h.page + 1), h.text, st, f"{h.rect[0]:.0f}, {h.rect[1]:.0f}"]
            for c, v in enumerate(vals):
                it = QTableWidgetItem(v)
                if h.changed:
                    bg = QColor(RED if h.ver == 1 else GREEN)
                    bg.setAlpha(55)
                    it.setBackground(bg)
                t.setItem(r, c, it)
        nchg = sum(1 for h in hits if h.changed)
        npg = len({h.page for h in hits})
        self.lbl_query.setText(f"Kết quả: <b>{len(hits)}</b> trên <b>{npg}</b> trang &nbsp;·&nbsp; "
                               f"<span style='color:{ORANGE.name()}'>có thay đổi: <b>{nchg}</b></span>")
        self.redraw_overlays()

    def on_query_click(self, item):
        h = self.query_hits[item.row()]
        self._goto(h.page, qrect(h.rect), h.ver, 8.0)

    def clear_query(self):
        self.query_hits = []
        self.q_table.setRowCount(0)
        self.lbl_query.setText("Kết quả: —")
        self.redraw_overlays()

    def export_query(self):
        if not self.query_hits:
            return
        rows = [[f"V{h.ver}", h.page + 1, h.text, "Thay đổi" if h.changed else "Không đổi"]
                + [round(v, 1) for v in h.rect] for h in self.query_hits]
        self.write_table("Query", ["Bản", "Trang", "Text", "Trạng thái", "x0", "y0", "x1", "y1"], rows,
                         "query_result.xlsx", status_col=3)

    # ========================================================= highlight ==
    def _set_hl_color(self, c: QColor):
        self.hl_color = QColor(c)
        for v in (self.viewA, self.viewB):
            if self.tool == "highlight":
                v.set_tool("highlight", self.hl_color)

    def _pick_hl_color(self):
        c = QColorDialog.getColor(self.hl_color, self, "Màu highlight")
        if c.isValid():
            for b in self.hl_group.buttons():
                b.setChecked(False)
            self._set_hl_color(c)

    def hl_add_from_rect(self, view: PdfView, r: QRectF):
        src = self.view_src.get(view, 1)
        ver = src if src in (1, 2) else 1
        doc = self.docs[ver]
        if not doc or not doc.has_page(self.page):
            return
        fr = frect(r)
        mode = self.cb_hl_mode.currentIndex()
        rects, text = ([], "")
        if mode != 1:
            rects, text = core.snap_text_rects(doc, self.page, fr)
        if rects:
            kind = "text"
            rects = [(x0 - 0.5, y0 - 0.5, x1 + 0.5, y1 + 0.5) for x0, y0, x1, y1 in rects]
        elif mode == 2:
            self.msg("Không có chữ trong vùng kéo (kiểu 'Chỉ tô chữ').")
            return
        else:
            kind = "area"
            rects = [ttuple(r)]
            text = "\n".join(doc.lines_in(self.page, fr))
        self._hl_id += 1
        h = core.Highlight(self._hl_id, ver, self.page, rects, self.hl_color.name(), self.sl_op.value() / 100,
                           kind, text, "", self.ed_author.text().strip(), datetime.now().strftime("%Y-%m-%d %H:%M"))
        self.highlights.append(h)
        self.hl_sel = h.id
        self._refresh_hl_table()
        self.redraw_overlays()

    def _refresh_hl_table(self):
        t = self.hl_table
        t.setRowCount(0)
        for h in self.highlights:
            r = t.rowCount()
            t.insertRow(r)
            vals = ["", f"PDF {h.ver}", str(h.page + 1), "Chữ" if h.kind == "text" else "Vùng",
                    h.text.replace("\n", " ⏎ "), h.comment.replace("\n", " "), h.author, h.date]
            for c, v in enumerate(vals):
                it = QTableWidgetItem(v)
                it.setToolTip(h.text if c == 4 else v)
                if c == 0:
                    it.setBackground(QColor(h.color))
                t.setItem(r, c, it)
            if h.id == self.hl_sel:
                t.selectRow(r)
        cnt = Counter(h.ver for h in self.highlights)
        self.lbl_hl.setText(f"<b>{len(self.highlights)}</b> highlight"
                            + "".join(f" &nbsp;·&nbsp; PDF {v}: <b>{n}</b>" for v, n in sorted(cnt.items())))

    def _hl_by_row(self, row: int):
        return self.highlights[row] if 0 <= row < len(self.highlights) else None

    def _hl_selected(self):
        for h in self.highlights:
            if h.id == self.hl_sel:
                return h
        return None

    def _on_hl_click(self, item):
        h = self._hl_by_row(item.row())
        if not h:
            return
        self.hl_sel = h.id
        if self.mode == "v2" and h.ver == 1 or self.mode == "v1" and h.ver == 2:
            self.set_mode(f"v{h.ver}")
        self._goto(h.page, qrect(h.bbox()), h.ver, 2.5)
        self.redraw_overlays()

    def _hl_at(self, view: PdfView, pos: QPointF):
        src = self.view_src.get(view, 1)
        for h in reversed(self.highlights):
            if h.page != self.page or (src in (1, 2) and h.ver != src):
                continue
            for r in h.rects:
                if self._to_view(qrect(r), h.ver, src).adjusted(-1, -1, 1, 1).contains(pos):
                    return h
        return None

    def on_view_context(self, view: PdfView, pos: QPointF, gpos):
        h = self._hl_at(view, pos)
        if not h:
            return
        self.hl_sel = h.id
        self._refresh_hl_table()
        self.redraw_overlays()
        m = QMenu(self)
        m.addAction("💬 Ghi chú…", lambda: self._hl_comment(h))
        cm = m.addMenu("🎨 Đổi màu")
        for hexc, name in HL_SWATCHES:
            cm.addAction(f"■ {name}", lambda c=hexc: self._hl_recolor(h, QColor(c)))
        om = m.addMenu("◐ Độ đậm")
        for op in (25, 45, 65, 85):
            om.addAction(f"{op}%", lambda o=op: self._hl_set_opacity(h, o / 100))
        if h.text:
            m.addAction("📋 Copy text", lambda: QApplication.clipboard().setText(h.text))
        m.addSeparator()
        m.addAction("🗑 Xoá", lambda: self._hl_remove(h))
        m.exec(gpos)

    def _hl_comment(self, h):
        if not h:
            return
        txt, ok = QInputDialog.getMultiLineText(self, "Ghi chú highlight",
                                                f"PDF {h.ver} · Trang {h.page + 1}\n{h.text[:120]}", h.comment)
        if ok:
            h.comment = txt.strip()
            self._refresh_hl_table()
            self.redraw_overlays()

    def _hl_recolor(self, h, c: QColor):
        if not h:
            return
        h.color = QColor(c).name()
        self._refresh_hl_table()
        self.redraw_overlays()

    def _hl_set_opacity(self, h, op: float):
        h.opacity = op
        self.redraw_overlays()

    def _hl_remove(self, h):
        if h in self.highlights:
            self.highlights.remove(h)
        if self.hl_sel == h.id:
            self.hl_sel = None
        self._refresh_hl_table()
        self.redraw_overlays()

    def hl_delete_selected(self):
        h = self._hl_selected()
        if h:
            self._hl_remove(h)

    def hl_undo(self):
        if self.highlights:
            self._hl_remove(self.highlights[-1])
            self.msg("Đã hoàn tác highlight cuối")

    def hl_clear(self):
        if self.highlights and QMessageBox.question(
                self, "Xoá hết", f"Xoá {len(self.highlights)} highlight?") == QMessageBox.Yes:
            self.highlights.clear()
            self.hl_sel = None
            self._refresh_hl_table()
            self.redraw_overlays()

    def hl_save_pdf(self):
        if not self.highlights:
            QMessageBox.information(self, "Lưu", "Chưa có highlight nào.")
            return
        saved = []
        for ver in (1, 2):
            items = [h for h in self.highlights if h.ver == ver]
            doc = self.docs[ver]
            if not items or not doc:
                continue
            stem = os.path.splitext(doc.path)[0]
            path, _ = QFileDialog.getSaveFileName(self, f"Lưu PDF {ver} kèm highlight", stem + "_markup.pdf",
                                                  "PDF (*.pdf)")
            if not path:
                continue
            if os.path.normcase(os.path.abspath(path)) == os.path.normcase(os.path.abspath(doc.path)):
                QMessageBox.warning(self, "Lưu", "File gốc đang mở, hãy lưu với tên khác (VD: …_markup.pdf).")
                continue
            try:
                with busy():
                    n = core.export_highlights(doc.path, path, items)
                saved.append(f"{path}  ({n} highlight)")
            except Exception as e:
                QMessageBox.critical(self, "Lưu", f"Lỗi: {e}")
        if saved and QMessageBox.question(self, "Đã lưu",
                                          "Highlight đã ghi thành annotation thật (mở được bằng Bluebeam / Acrobat):\n"
                                          + "\n".join(saved) + "\n\nMở file?") == QMessageBox.Yes:
            os.startfile(saved[0].split("  (")[0])

    def hl_export(self):
        if not self.highlights:
            return
        rows = [[f"PDF {h.ver}", h.page + 1, "Chữ" if h.kind == "text" else "Vùng", h.color, h.text, h.comment,
                 h.author, h.date] + [round(v, 1) for v in h.bbox()] for h in self.highlights]
        self.write_table("Markups", ["Bản", "Trang", "Kiểu", "Màu", "Text", "Ghi chú", "Người", "Ngày",
                                     "x0", "y0", "x1", "y1"], rows, "markups.xlsx")

    def hl_save_json(self):
        if not self.highlights:
            return
        path, _ = QFileDialog.getSaveFileName(self, "Lưu markup", os.path.join(self._last_dir, "markups.json"),
                                              "JSON (*.json)")
        if not path:
            return
        data = {"app": "PDF QA Viewer", "files": {str(v): (d.name if d else "") for v, d in self.docs.items()},
                "highlights": [h.to_json() for h in self.highlights]}
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        self.msg(f"Đã lưu markup: {path}")

    def hl_load_json(self):
        path, _ = QFileDialog.getOpenFileName(self, "Mở markup", self._last_dir, "JSON (*.json)")
        if not path:
            return
        try:
            with open(path, encoding="utf-8") as f:
                data = json.load(f)
            items = [core.Highlight.from_json(d) for d in data.get("highlights", [])]
        except Exception as e:
            QMessageBox.critical(self, "Markup", f"File không hợp lệ: {e}")
            return
        for h in items:
            self._hl_id += 1
            h.id = self._hl_id
        self.highlights.extend(items)
        self._refresh_hl_table()
        self.redraw_overlays()
        self.tabs.setCurrentIndex(4)

    # ============================================================ export ==
    def write_table(self, title, headers, rows, default_name, status_col=None):
        path, flt = QFileDialog.getSaveFileName(self, f"Xuất {title}", os.path.join(self._last_dir, default_name),
                                                "Excel (*.xlsx);;CSV (*.csv)")
        if not path:
            return
        try:
            if path.lower().endswith(".csv"):
                raise ImportError
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            wb = Workbook()
            ws = wb.active
            ws.title = title[:30]
            ws.append(headers)
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="2B4A8A")
                c.alignment = Alignment(vertical="center")
            fills = {"OK": "C6EFCE", "Không đổi": "C6EFCE", "CHANGED": "FFEB9C", "Thay đổi": "FFEB9C",
                     "DATA": "DDEBF7", "FAIL": "FFC7CE", "Text xoá": "FFC7CE", "Hình xoá": "FFC7CE", "Text thêm": "C6EFCE",
                     "Hình thêm": "C6EFCE", "Text sửa": "FFEB9C", "Hình sửa": "FFEB9C"}
            for row in rows:
                ws.append(row)
                if status_col is not None:
                    v = str(row[status_col])
                    if v in fills:
                        ws.cell(ws.max_row, status_col + 1).fill = PatternFill("solid", fgColor=fills[v])
            for i, h in enumerate(headers, 1):
                width = max([len(str(h))] + [min(60, len(str(r[i - 1]))) for r in rows[:500]]) + 2
                ws.column_dimensions[ws.cell(1, i).column_letter].width = width
            ws.freeze_panes = "A2"
            wb.save(path)
        except ImportError:
            if not path.lower().endswith(".csv"):
                path = os.path.splitext(path)[0] + ".csv"
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                w.writerow(headers)
                w.writerows(rows)
        except Exception as e:
            QMessageBox.critical(self, "Xuất", f"Lỗi: {e}")
            return
        self._last_dir = os.path.dirname(path)
        if QMessageBox.question(self, "Xuất xong", f"Đã lưu:\n{path}\n\nMở file?") == QMessageBox.Yes:
            os.startfile(path)
