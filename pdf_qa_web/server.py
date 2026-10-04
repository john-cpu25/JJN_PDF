"""
FastAPI backend for the PDF QA web version.

All heavy lifting (rendering, pixel diff, text diff, zone rules, crop, annotations)
re-uses ``pdf_qa.core`` unchanged, so results are identical to the desktop app.

Rectangles exchanged with the browser are display coordinates in PDF points
(origin top-left of the page as displayed), exactly like the desktop version.
"""
from __future__ import annotations

import io
import os
import shutil
import tempfile
import threading
import time
import uuid
import zipfile
from collections import OrderedDict
from contextlib import asynccontextmanager
from dataclasses import asdict
from typing import Any, Optional

import pymupdf as fitz
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from pdf_qa import core

# PyMuPDF is not thread safe -> serialise every fitz access.
LOCK = threading.RLock()
UPLOAD_DIR = os.path.join(tempfile.gettempdir(), "pdf_qa_web")
STATIC_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static")


# --------------------------------------------------------------------------- #
#  Document store (lazy open, LRU of open handles)
# --------------------------------------------------------------------------- #
class Store:
    def __init__(self, max_open: int = 10):
        self.meta: dict[str, dict] = {}            # id -> {path, name}
        self.open: OrderedDict[str, core.PdfDoc] = OrderedDict()
        self.max_open = max_open

    def add(self, path: str, name: str) -> str:
        did = uuid.uuid4().hex[:12]
        self.meta[did] = {"path": path, "name": name}
        return did

    def get(self, did: Optional[str]) -> Optional[core.PdfDoc]:
        if not did:
            return None
        if did in self.open:
            self.open.move_to_end(did)
            return self.open[did]
        m = self.meta.get(did)
        if m is None:
            raise HTTPException(404, f"Không tìm thấy tài liệu {did} (server đã khởi động lại?)")
        d = core.PdfDoc(m["path"])
        d.name = m["name"]
        self.open[did] = d
        while len(self.open) > self.max_open:
            _, old = self.open.popitem(last=False)
            old.close()
        return d

    def close(self, did: str):
        d = self.open.pop(did, None)
        if d:
            d.close()

    def remove(self, did: str):
        self.close(did)
        m = self.meta.pop(did, None)
        if m:
            try:
                os.remove(m["path"])
            except OSError:
                pass


store = Store()


class LRU(OrderedDict):
    def __init__(self, maxsize: int):
        super().__init__()
        self.maxsize = maxsize

    def get_or(self, key, fn):
        if key in self:
            self.move_to_end(key)
            return self[key]
        val = fn()
        self[key] = val
        while len(self) > self.maxsize:
            self.popitem(last=False)
        return val


_png_cache = LRU(12)
_pd_cache = LRU(4)
_td_cache = LRU(64)


# --------------------------------------------------------------------------- #
#  Helpers
# --------------------------------------------------------------------------- #
def png_bytes(arr) -> bytes:
    import numpy as np
    arr = np.ascontiguousarray(arr)
    h, w = arr.shape[:2]
    cs = fitz.csGRAY if arr.ndim == 2 else fitz.csRGB
    return fitz.Pixmap(cs, w, h, arr.tobytes(), False).tobytes("png")


def png_response(data: bytes) -> Response:
    return Response(data, media_type="image/png", headers={"Cache-Control": "private, max-age=600"})


def parse_clip(clip: Optional[str]) -> Optional[fitz.Rect]:
    if not clip:
        return None
    try:
        x0, y0, x1, y1 = [float(v) for v in clip.split(",")]
    except ValueError:
        raise HTTPException(400, "clip phải dạng x0,y0,x1,y1")
    return fitz.Rect(x0, y0, x1, y1)


def shifted(r, dx: float, dy: float) -> fitz.Rect:
    r = fitz.Rect(r)
    return fitz.Rect(r.x0 + dx, r.y0 + dy, r.x1 + dx, r.y1 + dy)


def rt(r) -> list:
    return [round(float(v), 2) for v in tuple(r)[:4]]


def need_page(d: Optional[core.PdfDoc], page: int) -> bool:
    return bool(d is not None and d.has_page(page))


def get_pdiff(d1id, d2id, page, dpi, thr, tol, dx, dy) -> core.PixelDiff:
    key = (d1id, d2id, page, round(dpi, 2), thr, tol, dx, dy)

    def _do():
        d1, d2 = store.get(d1id), store.get(d2id)
        z = dpi / 72.0
        return core.pixel_diff(d1.render(page, dpi, True), d2.render(page, dpi, True), dpi, thr, tol,
                               (int(round(-dx * z)), int(round(-dy * z))))
    return _pd_cache.get_or(key, _do)


def changed_keys(d1id, d2id, page, dx, dy):
    key = (d1id, d2id, page, dx, dy)

    def _do():
        td = core.text_diff(store.get(d1id).words(page), store.get(d2id).words(page), (dx, dy))
        k1 = {tuple(w[:4]) for w in td.removed} | {tuple(a[:4]) for a, _ in td.changed}
        k2 = {tuple(w[:4]) for w in td.added} | {tuple(b[:4]) for _, b in td.changed}
        return k1, k2
    return _td_cache.get_or(key, _do)


def zone_dict(z: core.Zone) -> dict:
    d = asdict(z)
    d["rect"] = rt(z.rect)
    return d


def file_response(data: bytes, filename: str, media: str) -> Response:
    from urllib.parse import quote
    return Response(data, media_type=media, headers={
        "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}",
        "X-Filename": quote(filename),
        "Access-Control-Expose-Headers": "X-Filename",
    })


def save_upload(f: UploadFile) -> tuple[str, str, str]:
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    name = os.path.basename(f.filename or "file.pdf")
    path = os.path.join(UPLOAD_DIR, f"{uuid.uuid4().hex}.pdf")
    with open(path, "wb") as out:
        shutil.copyfileobj(f.file, out, 1024 * 1024)
    return path, name, store.add(path, name)


def doc_info(did: str) -> dict:
    d = store.get(did)
    pages = []
    for i in range(d.page_count):
        r = d.page_rect(i)
        pages.append([round(r.width, 3), round(r.height, 3)])
    return {"id": did, "name": d.name, "page_count": d.page_count, "pages": pages}


def cleanup_uploads(max_age_h: float = 24):
    if not os.path.isdir(UPLOAD_DIR):
        return
    now = time.time()
    for fn in os.listdir(UPLOAD_DIR):
        p = os.path.join(UPLOAD_DIR, fn)
        try:
            if now - os.path.getmtime(p) > max_age_h * 3600:
                os.remove(p)
        except OSError:
            pass


# --------------------------------------------------------------------------- #
#  App
# --------------------------------------------------------------------------- #
@asynccontextmanager
async def _lifespan(_app):
    cleanup_uploads()
    yield
    with LOCK:
        for did in list(store.open):
            store.close(did)


app = FastAPI(title="JNN PDF Web", version="1.0.0", lifespan=_lifespan)


@app.exception_handler(Exception)
async def _err(_req, exc: Exception):
    return JSONResponse({"detail": f"{type(exc).__name__}: {exc}"}, status_code=500)


@app.get("/api/meta")
def meta():
    return {"rules": core.RULES, "labels": core.DiffItem.LABELS,
            "user": os.environ.get("USERNAME", os.environ.get("USER", ""))}


# ------------------------------------------------------------------ documents
@app.post("/api/docs")
def upload_doc(file: UploadFile = File(...)):
    path, name, did = save_upload(file)
    try:
        with LOCK:
            return doc_info(did)
    except HTTPException:
        raise
    except Exception as e:
        store.remove(did)
        raise HTTPException(400, f"Không mở được file '{name}': {e}")


@app.delete("/api/docs/{did}")
def delete_doc(did: str):
    with LOCK:
        store.remove(did)
    return {"ok": True}


@app.get("/api/docs/{did}")
def get_doc(did: str):
    with LOCK:
        return doc_info(did)


@app.get("/api/docs/{did}/render")
def render(did: str, page: int, dpi: float, clip: Optional[str] = None):
    dpi = max(10.0, min(1600.0, dpi))
    with LOCK:
        d = store.get(did)
        if not d.has_page(page):
            raise HTTPException(404, "Trang không tồn tại")
        c = parse_clip(clip)
        if c is None:
            data = _png_cache.get_or((did, page, round(dpi, 2)), lambda: png_bytes(d.render(page, dpi)))
        else:
            data = png_bytes(d.render(page, dpi, False, c))
    return png_response(data)


@app.get("/api/overlay")
def overlay(d1: str, d2: str, page: int, dpi: float, thr: int = 200, tol: int = 1,
            dx: float = 0, dy: float = 0, clip: Optional[str] = None):
    dpi = max(10.0, min(1600.0, dpi))
    with LOCK:
        D1, D2 = store.get(d1), store.get(d2)
        if not (D1.has_page(page) and D2.has_page(page)):
            raise HTTPException(404, "Cả 2 bản phải có trang này")
        c = parse_clip(clip)
        if c is None:
            key = ("ov", d1, d2, page, round(dpi, 2), thr, tol, dx, dy)
            data = _png_cache.get_or(key, lambda: png_bytes(
                get_pdiff(d1, d2, page, dpi, thr, tol, dx, dy).overlay_rgb()))
        else:
            g1 = D1.render(page, dpi, True, c)
            g2 = D2.render(page, dpi, True, shifted(c, dx, dy))
            data = png_bytes(core.pixel_diff(g1, g2, dpi, thr, tol).overlay_rgb())
    return png_response(data)


# ------------------------------------------------------------------- compare
class Settings(BaseModel):
    d1: Optional[str] = None
    d2: Optional[str] = None
    thr: int = 200
    tol: int = 1
    dx: float = 0
    dy: float = 0


class CompareReq(Settings):
    page: int
    dpi: float = 100
    pixel: bool = True
    text: bool = True
    use_cache: bool = True


@app.post("/api/compare")
def compare(req: CompareReq):
    with LOCK:
        D1, D2 = store.get(req.d1), store.get(req.d2)
        if not (need_page(D1, req.page) and need_page(D2, req.page)):
            return {"page": req.page, "items": [], "missing": True}
        pd = None
        if req.pixel and req.use_cache:
            pd = get_pdiff(req.d1, req.d2, req.page, req.dpi, req.thr, req.tol, req.dx, req.dy)
        items, _ = core.compare_page(D1, D2, req.page, req.dpi, req.thr, req.tol, (req.dx, req.dy),
                                     req.pixel, req.text, pdiff=pd)
    return {"page": req.page, "items": [
        {"kind": it.kind, "label": it.label, "page": it.page, "rect": rt(it.rect), "text": it.text}
        for it in items]}


# ----------------------------------------------------------------- read data
class RectReq(Settings):
    page: int
    rect: list[float]


@app.post("/api/read")
def read_region(req: RectReq):
    import base64
    from collections import Counter
    with LOCK:
        D1, D2 = store.get(req.d1), store.get(req.d2)
        p = req.page
        f1 = fitz.Rect(req.rect)
        f2 = shifted(f1, req.dx, req.dy)
        has1, has2 = need_page(D1, p), need_page(D2, p)
        lines1 = D1.lines_in(p, f1) if has1 else None
        lines2 = D2.lines_in(p, f2) if has2 else None
        td = None
        if has1 and has2:
            rows = core.line_diff(lines1, lines2)
            t = core.text_diff(D1.words_in(p, f1), D2.words_in(p, f2), (req.dx, req.dy))
            td = {"removed": [rt(w) for w in t.removed], "added": [rt(w) for w in t.added],
                  "changed": [[rt(a), rt(b)] for a, b in t.changed]}
        else:
            rows = [("equal", l, "") for l in (lines1 or [])] + [("equal", "", l) for l in (lines2 or [])]
        long = max(f1.width, f1.height)
        pdpi = max(72.0, min(600.0, 900 * 72 / max(long, 1)))
        px, preview = None, None
        try:
            if has1 and has2:
                pd = core.pixel_diff(D1.render(p, pdpi, True, f1), D2.render(p, pdpi, True, f2), pdpi,
                                     req.thr, req.tol)
                arr = pd.overlay_rgb()
                px = {"removed": pd.n_removed, "added": pd.n_added}
            else:
                d, f = (D1, f1) if has1 else (D2, f2)
                arr = d.render(p, pdpi, False, f) if d else None
            if arr is not None:
                preview = "data:image/png;base64," + base64.b64encode(png_bytes(arr)).decode()
        except Exception:
            preview = None
    cnt = Counter(st for st, _, _ in rows)
    return {"page": p, "rect": rt(f1), "has1": has1, "has2": has2,
            "n1": None if lines1 is None else len(lines1), "n2": None if lines2 is None else len(lines2),
            "rows": [list(r) for r in rows], "count": dict(cnt), "td": td, "px": px, "preview": preview}


class TextInReq(BaseModel):
    doc: str
    page: int
    rect: list[float]


@app.post("/api/text_in")
def text_in(req: TextInReq):
    with LOCK:
        d = store.get(req.doc)
        if not d.has_page(req.page):
            return {"lines": []}
        return {"lines": d.lines_in(req.page, fitz.Rect(req.rect))}


# --------------------------------------------------------------------- zones
class ZoneIn(BaseModel):
    name: str
    page: int
    rect: list[float]
    rule: str = "compare"
    expected: str = ""


class ZoneReq(Settings):
    zones: list[ZoneIn]


@app.post("/api/zones/check")
def zones_check(req: ZoneReq):
    out = []
    with LOCK:
        D1, D2 = store.get(req.d1), store.get(req.d2)
        for zi in req.zones:
            z = core.Zone(zi.name, zi.page, tuple(zi.rect), zi.rule, zi.expected)
            core.check_zone(z, D1, D2, (req.dx, req.dy), 200, req.thr, req.tol)
            out.append(zone_dict(z))
    return {"zones": out}


@app.post("/api/batch/file")
def batch_file(file: UploadFile = File(...), templates: str = Form(...)):
    import json
    temps = [ZoneIn(**t) for t in json.loads(templates)]
    path, name, did = save_upload(file)
    results = []
    with LOCK:
        try:
            doc = store.get(did)
        except Exception as e:
            z = core.Zone("(file)", 0, (0, 0, 0, 0), "extract")
            z.status, z.detail = "FAIL", f"Không mở được: {e}"
            store.remove(did)
            return {"doc": None, "name": name, "results": [zone_dict(z)]}
        for p in range(doc.page_count):
            pr = doc.page_rect(p)
            for t in temps:
                rule = "extract" if t.rule == "compare" else t.rule
                z = core.Zone(t.name, p, tuple(t.rect), rule, t.expected)
                if not fitz.Rect(t.rect).intersects(pr):
                    z.status, z.detail = "N/A", "vùng nằm ngoài trang (khác khổ giấy?)"
                else:
                    core.check_zone(z, doc, None)
                results.append(zone_dict(z))
        store.close(did)
    return {"doc": did, "name": name, "results": results}


# --------------------------------------------------------------------- query
class QueryReq(Settings):
    pattern: str
    vers: list[int] = [1, 2]
    pages: list[int]
    regex: bool = False
    case: bool = False
    rects: Optional[dict[int, list[list[float]]]] = None
    changed_only: bool = False


@app.post("/api/query")
def query(req: QueryReq):
    hits = []
    with LOCK:
        docs = {1: store.get(req.d1), 2: store.get(req.d2)}
        for v in req.vers:
            d = docs.get(v)
            if not d:
                continue
            rr = req.rects
            if rr is not None and v == 2:
                rr = {p: [(r[0] + req.dx, r[1] + req.dy, r[2] + req.dx, r[3] + req.dy) for r in lst]
                      for p, lst in rr.items()}
            hits += core.query_words(d, v, req.pages, req.pattern, req.regex, req.case, rr)
        for h in hits:
            if need_page(docs[1], h.page) and need_page(docs[2], h.page):
                k1, k2 = changed_keys(req.d1, req.d2, h.page, req.dx, req.dy)
                h.changed = tuple(h.rect) in (k1 if h.ver == 1 else k2)
    if req.changed_only:
        hits = [h for h in hits if h.changed]
    hits.sort(key=lambda h: (h.page, h.ver, h.rect[1], h.rect[0]))
    return {"hits": [{"ver": h.ver, "page": h.page, "rect": rt(h.rect), "text": h.text, "changed": h.changed}
                     for h in hits]}


# ---------------------------------------------------------------- highlights
class SnapReq(BaseModel):
    doc: str
    page: int
    rect: list[float]
    mode: int = 0        # 0 auto, 1 area only, 2 text only


@app.post("/api/snap")
def snap(req: SnapReq):
    with LOCK:
        d = store.get(req.doc)
        if not d.has_page(req.page):
            return {"kind": None, "rects": [], "text": ""}
        fr = fitz.Rect(req.rect)
        rects, text = ([], "")
        if req.mode != 1:
            rects, text = core.snap_text_rects(d, req.page, fr)
        if rects:
            rects = [[round(x0 - 0.5, 2), round(y0 - 0.5, 2), round(x1 + 0.5, 2), round(y1 + 0.5, 2)]
                     for x0, y0, x1, y1 in rects]
            return {"kind": "text", "rects": rects, "text": text}
        if req.mode == 2:
            return {"kind": None, "rects": [], "text": ""}
        return {"kind": "area", "rects": [rt(fr)], "text": "\n".join(d.lines_in(req.page, fr))}


class HighlightExportReq(BaseModel):
    doc: str
    items: list[dict[str, Any]]


@app.post("/api/highlights/export")
def highlights_export(req: HighlightExportReq):
    items = [core.Highlight.from_json(h) for h in req.items]
    with LOCK:
        d = store.get(req.doc)
        tmp = os.path.join(UPLOAD_DIR, f"out_{uuid.uuid4().hex}.pdf")
        try:
            core.export_highlights(d.path, tmp, items)
            with open(tmp, "rb") as f:
                data = f.read()
        finally:
            try:
                os.remove(tmp)
            except OSError:
                pass
    stem = os.path.splitext(d.name)[0]
    return file_response(data, f"{stem}_markup.pdf", "application/pdf")


# ---------------------------------------------------------------------- crop
class CropRegion(BaseModel):
    page: int
    rect: list[float]
    name: str = ""


class CropReq(Settings):
    regions: list[CropRegion]
    use1: bool = True
    use2: bool = False
    use_diff: bool = False
    as_pdf: bool = True
    dpi: int = 300


@app.post("/api/crop")
def crop(req: CropReq):
    dpi = max(72, min(1200, req.dpi))
    outs: list[tuple[str, bytes]] = []
    with LOCK:
        D1, D2 = store.get(req.d1), store.get(req.d2)
        base = f"crop_p{req.regions[0].page + 1}" if req.regions else "crop"
        items = []
        for reg in req.regions:
            p = reg.page
            f1 = fitz.Rect(reg.rect)
            f2 = shifted(f1, req.dx, req.dy)
            sfx = f"_{reg.name}" if reg.name else ""
            if req.use1 and need_page(D1, p):
                items.append((D1, p, f1, f"{sfx}_v1"))
            if req.use2 and need_page(D2, p):
                items.append((D2, p, f2, f"{sfx}_v2"))
            if req.use_diff and need_page(D1, p) and need_page(D2, p):
                arr = core.pixel_diff(D1.render(p, dpi, True, f1), D2.render(p, dpi, True, f2), dpi,
                                      req.thr, req.tol).overlay_rgb()
                outs.append((f"{base}{sfx}_diff.png", png_bytes(arr)))
        if req.as_pdf and items:
            tmp = os.path.join(UPLOAD_DIR, f"crop_{uuid.uuid4().hex}.pdf")
            os.makedirs(UPLOAD_DIR, exist_ok=True)
            try:
                core.export_crop_pdf(tmp, [(d, p, f) for d, p, f, _ in items])
                with open(tmp, "rb") as fh:
                    outs.insert(0, (f"{base}.pdf", fh.read()))
            finally:
                try:
                    os.remove(tmp)
                except OSError:
                    pass
        elif not req.as_pdf:
            for d, p, f, sfx in items:
                outs.append((f"{base}{sfx}.png", d.page(p).get_pixmap(dpi=dpi, clip=f, alpha=False).tobytes("png")))
    if not outs:
        raise HTTPException(400, "Không có gì để xuất (vùng crop rỗng?)")
    if len(outs) == 1:
        name, data = outs[0]
        return file_response(data, name, "application/pdf" if name.endswith(".pdf") else "image/png")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in outs:
            zf.writestr(name, data)
    return file_response(buf.getvalue(), f"{base}.zip", "application/zip")


# --------------------------------------------------------------------- excel
FILLS = {"OK": "C6EFCE", "Không đổi": "C6EFCE", "CHANGED": "FFEB9C", "Thay đổi": "FFEB9C",
         "DATA": "DDEBF7", "FAIL": "FFC7CE", "N/A": "EDEDED", "Text xoá": "FFC7CE", "Hình xoá": "FFC7CE",
         "Text thêm": "C6EFCE", "Hình thêm": "C6EFCE", "Text sửa": "FFEB9C", "Hình sửa": "FFEB9C"}


class Sheet(BaseModel):
    title: str
    headers: list[str]
    rows: list[list[Any]]
    status_col: Optional[int] = None
    status_fallback: Optional[str] = None     # fill for status values not in FILLS


class ExcelReq(BaseModel):
    filename: str = "export.xlsx"
    sheets: list[Sheet]


@app.post("/api/excel")
def excel(req: ExcelReq):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    for sh in req.sheets:
        ws = wb.create_sheet(sh.title[:30] or "Sheet")
        ws.append(sh.headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2B4A8A")
            c.alignment = Alignment(vertical="center")
        for row in sh.rows:
            ws.append(row)
            if sh.status_col is not None and sh.status_col < len(row):
                v = str(row[sh.status_col])
                col = FILLS.get(v, sh.status_fallback)
                if col:
                    ws.cell(ws.max_row, sh.status_col + 1).fill = PatternFill("solid", fgColor=col)
        for i, h in enumerate(sh.headers, 1):
            width = max([len(str(h))] + [min(60, len(str(r[i - 1]))) for r in sh.rows[:500] if len(r) >= i]) + 2
            ws.column_dimensions[ws.cell(1, i).column_letter].width = width
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return file_response(buf.getvalue(), req.filename,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------------------------------------------------------------- static UI
app.mount("/", StaticFiles(directory=STATIC_DIR, html=True), name="static")
